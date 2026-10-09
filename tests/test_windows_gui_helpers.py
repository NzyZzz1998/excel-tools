"""Exercise the real Tk window; business fixtures stay in temporary directories."""

import gc
import queue
import tempfile
import threading
import time
import tkinter as tk
from tkinter import font as tkfont
import unittest
import weakref
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from excel_unmerge_gui import APP_VERSION, Application
from excel_unmerge_fill import process_file


class WindowTests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.attributes("-alpha", 0)
        self.app = Application(self.root)
        self.root.update()
        self.directory = tempfile.TemporaryDirectory(prefix="excel-gui-test-")
        self.folder = Path(self.directory.name)

    def tearDown(self):
        # destroy() removes Tcl widgets, but Python cycles (including mocked
        # exception tracebacks) may still own the interpreter and its variables.
        # Release them on their creator thread before another test starts workers.
        self.assertIs(threading.current_thread(), threading.main_thread())
        app_ref, root_ref = weakref.ref(self.app), weakref.ref(self.app.root)
        try:
            if self.root is not None:
                self.root.destroy()
        finally:
            self.app = None
            self.root = None
            gc.collect()
            self.directory.cleanup()
        self.assertIsNone(app_ref(), "The test retained its destroyed Tk application")
        self.assertIsNone(root_ref(), "The test retained its destroyed Tcl interpreter")

    def workbook(self, name, merged):
        path = self.folder / name
        book = Workbook()
        book.active["A1"] = "测试内容"
        if merged:
            book.active.merge_cells("A1:A3")
        book.save(path)
        book.close()
        return path

    def finish(self):
        deadline = time.monotonic() + 10
        while self.app.running and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)
        self.assertFalse(self.app.running, "Batch did not finish")

    def wait_for(self, condition):
        deadline = time.monotonic() + 5
        while not condition() and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)
        self.assertTrue(condition(), "Window did not reach the expected state")

    def test_elapsed_and_real_phases_do_not_advance_file_completion(self):
        source = str(self.folder / "慢速合成.xlsx")
        commands = queue.Queue()
        self.app.files = (source,)

        def slow_process(file, all_merges, progress=None):
            progress(dict(phase="reading", sheet="测试表", completed=1200,
                          total=None, unit="rows"))
            while True:
                detail = commands.get(timeout=5)
                if detail is None:
                    return None, []
                progress(detail)

        with patch("excel_unmerge_gui.process_file", side_effect=slow_process):
            self.app.start()
            try:
                self.wait_for(lambda: "1,200 行" in self.app.status.get())
                self.assertIn("读取", self.app.status.get())
                self.assertIn("测试表", self.app.status.get())
                self.assertNotIn("%", self.app.status.get())
                self.assertNotIn("1,200 行 /", self.app.status.get())
                self.assertEqual(self.app.completed, 0)
                self.assertEqual(float(self.app.progress["value"]), 0)
                self.assertEqual(float(self.app.progress["maximum"]), 1)

                # Advance the per-file origin, then let the real 100ms Tk poll
                # refresh elapsed time without any further worker notification.
                self.app.current_file_started_at -= 65
                self.wait_for(lambda: "已用时 01:" in self.app.status.get())
                commands.put(dict(phase="filling", sheet="测试表", completed=1200,
                                  total=3000, unit="rows"))
                self.wait_for(lambda: "填充" in self.app.status.get())
                self.assertIn("1,200 行 / 3,000 行", self.app.status.get())
                commands.put(dict(phase="saving", sheet=None, completed=4096,
                                  total=8192, unit="bytes"))
                self.wait_for(lambda: "保存" in self.app.status.get())
                self.assertIn("4,096 字节 / 8,192 字节", self.app.status.get())
                self.assertEqual(float(self.app.progress["value"]), 0)

                self.app.close()
                stopped_status = self.app.status.get()
                commands.put(dict(phase="saving", sheet=None, completed=8192,
                                  total=8192, unit="bytes"))
                self.wait_for(lambda: self.app.current_detail.get("completed") == 8192)
                self.assertEqual(self.app.status.get(), stopped_status)
                self.assertTrue(self.root.winfo_exists())
                self.assertTrue(self.app.running)
                self.assertEqual(self.app.completed, 0)
            finally:
                commands.put(None)
                self.finish()
        final_status = self.app.status.get()
        self.app.refresh_running_status()
        self.assertEqual(self.app.status.get(), final_status)
        self.assertIsNone(self.app.current_file_started_at)
        self.assertIn("上次任务已处理 1/1", final_status)
        self.assertNotIn("已用时", final_status)
        self.assertEqual(float(self.app.progress["value"]), 1)
        self.assertEqual(APP_VERSION, "1.1.1")
        self.assertIn("v1.1.1", self.root.title())

    def test_elapsed_restarts_for_each_file(self):
        first, later = "合成第一份.xlsx", "合成第二份.xlsx"
        self.app.files = self.app.run_files = (first, later)
        self.app.file_states = {first: "pending", later: "pending"}
        self.app.running = True
        self.app.events.put(("started", first, 1, 2))
        self.app.poll_results()
        self.app.current_file_started_at -= 65
        self.app.refresh_running_status()
        self.assertIn("已用时 01:", self.app.status.get())
        self.app.events.put(("result", first, None, [], None))
        self.app.events.put(("started", later, 2, 2))
        self.app.poll_results()
        self.assertIn(later, self.app.status.get())
        self.assertIn("已用时 00:00", self.app.status.get())
        self.assertIn("本轮第 2/2 个，已完成 1 个", self.app.status.get())
        self.assertEqual(float(self.app.progress["value"]), 1)
        self.app.events.put(("result", later, None, [], None))
        self.app.events.put(("done", 2, 0, ()))
        self.app.poll_results()
        self.assertFalse(self.app.running)

    def test_progress_after_error_cannot_overwrite_failure_or_done_status(self):
        source = str(self.folder / "合成失败.xlsx")
        self.app.files = (source,)
        self.app.run_files = (source,)
        self.app.file_states = {source: "pending"}
        self.app.running = True
        self.app.events.put(("started", source, 1, 1))
        self.app.events.put(("progress", source, dict(
            phase="reading", sheet=None, completed=10, total=None, unit="rows")))
        self.app.events.put(("result", source, None, [], "合成错误"))
        self.app.events.put(("progress", source, dict(
            phase="saving", sheet=None, completed=20, total=20, unit="bytes")))
        self.app.poll_results()
        failed_status = self.app.status.get()
        self.assertIn("失败", failed_status)
        self.assertNotIn("已用时", failed_status)
        self.assertIsNone(self.app.current_file_started_at)
        self.app.refresh_running_status()
        self.assertEqual(self.app.status.get(), failed_status)
        self.app.events.put(("done", 1, 1, ()))
        self.app.poll_results()
        done_status = self.app.status.get()
        self.app.events.put(("progress", source, dict(
            phase="reading", sheet=None, completed=30, total=None, unit="rows")))
        self.app.poll_results()
        self.assertEqual(self.app.status.get(), done_status)
        self.assertIsNone(self.app.current_detail)
        self.assertEqual(self.app.failed_files, (source,))

    def check_recovery_displays_old_target_and_keeps_next_selection(self, state):
        old = str(self.folder / "上次任务A.xlsx")
        new = str(self.folder / "待开始B.xlsx")
        self.app.files = (old,)
        self.app.file_states = {old: state}
        self.app.batch_all_merges = False
        self.app.update_counts()
        with patch("excel_unmerge_gui.filedialog.askopenfilenames", return_value=(new,)):
            self.app.choose()
        self.assertEqual(self.app.files, (new,))
        self.assertIn(new, self.app.file_list.get("1.0", "end"))
        self.assertIn("待开始", self.app.input_frame.cget("text"))
        self.assertIn("恢复按钮仍处理上次任务", self.app.status.get())
        button = self.app.retry_button if state == "failed" else self.app.resume_button
        self.assertIn("上次", button.cget("text"))
        self.assertIn("（1）", button.cget("text"))
        self.app.all_merges.set(True)
        entered = threading.Event()
        release = threading.Event()
        processed = []

        def process(file, all_merges, progress=None):
            processed.append((file, all_merges))
            if file == old:
                entered.set()
                if not release.wait(5):
                    raise AssertionError("Recovery was not released")
            return None, []

        with patch("excel_unmerge_gui.process_file", side_effect=process):
            if state == "failed":
                self.app.retry_failed()
            else:
                self.app.resume_pending()
            try:
                self.assertTrue(entered.wait(5))
                self.wait_for(lambda: self.app.current_file == old)
                self.assertEqual(self.app.files, (new,))
                self.assertEqual(self.app.file_list.get("1.0", "end-1c"), old)
                self.assertIn("上次任务", self.app.input_frame.cget("text"))
                self.assertIn(Path(old).name, self.app.status.get())
                self.assertNotIn(Path(new).name, self.app.status.get())
                self.assertIs(self.app.all_merges.get(), False)
            finally:
                release.set()
                self.finish()
            self.assertEqual(self.app.files, (new,))
            self.assertEqual(self.app.file_list.get("1.0", "end-1c"), new)
            self.assertIn("待开始", self.app.input_frame.cget("text"))
            self.assertIn("上次任务已处理 1/1", self.app.status.get())
            self.assertIn("待开始选择已保留（1 个）", self.app.status.get())
            self.assertIn(old, self.app.results.get("1.0", "end"))
            self.assertNotIn(new, self.app.results.get("1.0", "end"))
            self.app.all_merges.set(True)
            self.app.start()
            self.finish()
        self.assertEqual(processed, [(old, False), (new, True)])
        self.assertEqual(tuple(self.app.file_states), (new,))
        self.assertIn(new, self.app.results.get("1.0", "end"))
        self.assertNotIn(old, self.app.results.get("1.0", "end"))

    def test_retry_shows_old_task_then_starts_preserved_new_selection(self):
        self.check_recovery_displays_old_target_and_keeps_next_selection("failed")

    def test_resume_shows_old_task_then_starts_preserved_new_selection(self):
        self.check_recovery_displays_old_target_and_keeps_next_selection("pending")

    def test_result_directory_unavailable_until_output_exists(self):
        button = getattr(self.app, "open_button", None)
        self.assertIsNotNone(button, "Results need an open-directory action")
        self.assertIn("disabled", button.state())
        for control in (self.app.choose_button, self.app.file_list, self.app.option,
                        self.app.start_button, self.app.progress, self.app.results,
                        self.app.retry_button, self.app.resume_button,
                        self.app.stop_button, button):
            self.assertGreater(control.winfo_width(), 1)
            self.assertGreater(control.winfo_height(), 1)
        for button in (self.app.retry_button, self.app.resume_button, self.app.stop_button):
            self.assertIn("disabled", button.state())

    def test_batch_reports_completed_files_and_enables_output_directory(self):
        successful = self.workbook("合并 数据.xlsx", True)
        unchanged = self.workbook("无需处理.xlsx", False)
        missing = self.folder / "不存在.xlsx"
        self.app.files = tuple(map(str, (successful, unchanged, missing)))
        self.app.start()
        self.assertEqual(str(self.app.progress["mode"]), "determinate")
        deadline = time.monotonic() + 10
        while self.app.running and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)
        self.assertFalse(self.app.running, "Batch did not finish")
        self.assertEqual(float(self.app.progress["value"]), 3)
        self.assertEqual(float(self.app.progress["maximum"]), 3)
        self.assertEqual(self.app.completed, 3)
        self.assertEqual(self.app.succeeded, 1)
        self.assertEqual(self.app.unchanged, 1)
        self.assertEqual(self.app.failed, 1)
        self.assertEqual(self.app.output_directory, self.folder.resolve())
        self.assertNotIn("disabled", self.app.open_button.state())
        output = successful.with_name("合并 数据_拆分填充.xlsx").resolve()
        self.assertTrue(output.is_file())
        log = self.app.results.get("1.0", "end")
        self.assertIn(str(output), log)
        self.assertIn(str(missing), log)
        self.assertIn(str(unchanged), log)

    def test_retry_only_failed_files_keeps_history_output_and_initial_options(self):
        successful = self.workbook("成功.xlsx", True)
        repaired = self.folder / "待修复.xlsx"
        repaired.write_bytes(b"not an xlsx")
        self.app.files = tuple(map(str, (successful, repaired)))
        self.app.start()
        self.finish()
        previous_log = self.app.results.get("1.0", "end-1c")
        self.assertIn("BadZipFile", previous_log)
        self.assertEqual(self.app.failed_files, (str(repaired),))
        self.assertNotIn("disabled", self.app.retry_button.state())

        book = Workbook()
        book.active["A1"] = "只保留横向"
        book.active.merge_cells("A1:B1")
        book.save(repaired)
        book.close()
        # Changing the next normal run's option must not change a retry's rule.
        self.app.all_merges.set(True)
        self.app.retry_failed()
        self.assertIs(self.app.all_merges.get(), False)
        self.finish()
        self.assertEqual(self.app.failed_files, ())
        self.assertEqual(self.app.succeeded, 1)
        self.assertEqual(self.app.unchanged, 1)
        self.assertEqual(self.app.completed, 2)
        self.assertEqual(self.app.output_directory, self.folder.resolve())
        self.assertEqual(len(list(self.folder.glob("成功_拆分填充*.xlsx"))), 1)
        self.assertFalse(repaired.with_name("待修复_拆分填充.xlsx").exists())
        self.assertTrue(self.app.results.get("1.0", "end").startswith(previous_log))
        self.assertIn("disabled", self.app.retry_button.state())

    def test_stop_and_continue_keep_current_file_safe_and_report_pending(self):
        first = self.workbook("第一份.xlsx", True)
        later = self.workbook("第二份.xlsx", True)
        self.app.files = tuple(map(str, (first, later)))
        entered = threading.Event()
        release = threading.Event()

        def delayed_process(file, all_merges, progress=None):
            if file == str(first):
                entered.set()
                if not release.wait(5):
                    raise AssertionError("Current file was not released")
            return process_file(file, all_merges=all_merges, progress=progress)

        with patch("excel_unmerge_gui.process_file", side_effect=delayed_process) as mocked:
            self.app.start()
            try:
                self.assertTrue(entered.wait(5))
                deadline = time.monotonic() + 2
                while self.app.current_file is None and time.monotonic() < deadline:
                    self.root.update()
                    time.sleep(.01)
                self.assertEqual(self.app.current_file, str(first))
                self.assertIn(first.name, self.app.status.get())
                self.app.close()
                self.assertTrue(self.app.running)
                self.assertTrue(self.root.winfo_exists())
                self.assertTrue(self.app.stop_event.is_set())
            finally:
                release.set()
                self.finish()
            self.assertEqual(mocked.call_count, 1)
            self.assertEqual(self.app.completed, 1)
            self.assertEqual(self.app.pending_files, (str(later),))
            self.assertIn("已停止", self.app.status.get())
            self.assertIn("待处理 1", self.app.status.get())
            self.assertNotIn("disabled", self.app.resume_button.state())
            self.assertFalse(later.with_name("第二份_拆分填充.xlsx").exists())
            self.app.all_merges.set(True)
            self.app.resume_pending()
            self.finish()
            self.assertEqual(mocked.call_count, 2)
            self.assertIs(mocked.call_args.kwargs["all_merges"], False)
        self.assertEqual(self.app.completed, 2)
        self.assertEqual(self.app.succeeded, 2)
        self.assertEqual(self.app.pending_files, ())
        self.assertIn("disabled", self.app.resume_button.state())

    def test_failure_and_unprocessed_files_have_independent_recovery_buttons(self):
        first = self.folder / "失败.xlsx"
        later = self.workbook("继续.xlsx", True)
        self.app.files = tuple(map(str, (first, later)))

        def fail_and_stop(file, all_merges, progress=None):
            self.app.stop_event.set()
            raise PermissionError("test access denied")

        with patch("excel_unmerge_gui.process_file", side_effect=fail_and_stop):
            self.app.start()
            self.finish()
        self.assertEqual(self.app.failed_files, (str(first),))
        self.assertEqual(self.app.pending_files, (str(later),))
        self.assertNotIn("disabled", self.app.retry_button.state())
        self.assertNotIn("disabled", self.app.resume_button.state())
        self.assertIn("失败 1", self.app.status.get())
        self.assertIn("待处理 1", self.app.status.get())
        self.workbook(first.name, True)
        self.app.retry_failed()
        self.finish()
        self.assertEqual(self.app.failed_files, ())
        self.assertEqual(self.app.pending_files, (str(later),))
        self.assertIn("待处理 1", self.app.status.get())
        self.assertNotIn("disabled", self.app.resume_button.state())

    def check_thread_memory_failure_recovery(self, target, stage):
        source = self.workbook("线程内存-" + stage + ".xlsx", True)
        original = source.read_bytes()
        callback_errors = []
        self.root.report_callback_exception = lambda *args: callback_errors.append(args[0].__name__)
        self.app.files = (str(source),)
        with patch(target, side_effect=MemoryError("synthetic thread allocation failure")):
            self.root.after(0, self.app.start)
            self.root.update()
        self.assertEqual(callback_errors, [])
        self.finish()
        self.assertEqual(self.app.pending_files, (str(source),))
        self.assertEqual(self.app.failed_files, ())
        self.assertEqual((self.app.completed, self.app.succeeded, self.app.failed), (0, 0, 0))
        self.assertEqual(list(self.folder.glob("*_拆分填充*.xlsx")), [])
        self.assertEqual(source.read_bytes(), original)
        self.assertNotIn("disabled", self.app.start_button.state())
        self.assertNotIn("disabled", self.app.resume_button.state())
        self.assertIn("disabled", self.app.stop_button.state())
        self.assertIn("内存", self.app.results.get("1.0", "end"))
        self.assertIn("MemoryError", self.app.results.get("1.0", "end"))

        self.app.resume_pending()
        self.finish()
        self.assertEqual(self.app.succeeded, 1)
        self.assertEqual(self.app.pending_files, ())
        fresh = self.workbook("新批次-" + stage + ".xlsx", True)
        self.app.files = (str(fresh),)
        self.app.start()
        self.finish()
        self.assertEqual(self.app.succeeded, 1)
        self.assertEqual(tuple(self.app.file_states), (str(fresh),))
        self.assertEqual(callback_errors, [])
        # Exercise a real idle-window close, then leave teardown only the files.
        with patch.object(self.root, "destroy", wraps=self.root.destroy) as destroy:
            self.app.close()
            destroy.assert_called_once_with()
        self.root = None

    def test_thread_construction_memory_failure_can_resume_and_close(self):
        self.check_thread_memory_failure_recovery("excel_unmerge_gui.threading.Thread", "构造")

    def test_thread_start_memory_failure_can_resume_and_close(self):
        self.check_thread_memory_failure_recovery("excel_unmerge_gui.threading.Thread.start", "启动")

    def test_new_controls_fit_minimum_window_width(self):
        self.app.file_states = {"failed-{}.xlsx".format(index): "failed" for index in range(100)}
        self.app.file_states.update({"pending-{}.xlsx".format(index): "pending" for index in range(100)})
        self.app.update_counts()
        self.root.geometry("{}x{}".format(*self.root.minsize()))
        self.app.status.set("正在处理：" + "合成明细" * 10
                            + ".xlsx（本轮第 1/100 个，已完成 0 个）\n"
                            "填充 · 工作表：合成测试 · 已处理 100,000 行 / 757,636 行 · 已用时 06:46")
        self.root.update()
        for button in (self.app.start_button, self.app.open_button,
                       self.app.retry_button, self.app.resume_button, self.app.stop_button):
            self.assertTrue(button.winfo_ismapped())
            self.assertGreaterEqual(button.winfo_width(), button.winfo_reqwidth())
            self.assertLessEqual(button.winfo_rootx() + button.winfo_width(),
                                 self.root.winfo_rootx() + self.root.winfo_width())

    def test_long_status_wraps_to_available_width_after_resizing(self):
        wide_size = (self.root.winfo_width(), self.root.winfo_height())
        narrow_size = self.root.minsize()
        widths = []
        for size in (wide_size, narrow_size):
            self.root.geometry("{}x{}".format(*size))
            self.root.update()
            widths.append(self.app.status_label.winfo_width())
        self.assertGreater(widths[0], widths[1], "The test must use distinct window widths")
        # Physical sizes and font metrics vary with DPI. Choose text between
        # two narrow lines and two wide lines so narrowing must add a line.
        label_font = tkfont.nametofont("TkDefaultFont", root=self.root)
        message = "正在处理："
        suffix = ".xlsx（本轮第 1/3 个，已处理 0 个）。"
        while label_font.measure(message + suffix) < sum(widths) - 20:
            message += "明细"
        self.app.status.set(message + suffix)
        heights = []
        for size in (wide_size, narrow_size, wide_size):
            self.root.geometry("{}x{}".format(*size))
            self.root.update()
            label = self.app.status_label
            self.assertLessEqual(label.winfo_reqwidth(), label.winfo_width())
            self.assertLessEqual(label.winfo_reqheight(), label.winfo_height())
            self.assertLessEqual(label.winfo_rootx() + label.winfo_width(),
                                 self.root.winfo_rootx() + self.root.winfo_width())
            heights.append(label.winfo_height())
        self.assertGreater(heights[1], heights[0], "A narrower status should use more lines")
        self.assertEqual(heights[2], heights[0])


if __name__ == "__main__":
    unittest.main()
