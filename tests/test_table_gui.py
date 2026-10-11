"""Real, hidden Toplevel tests; no user workbook or existing window is touched."""

import gc
import queue
import tempfile
import threading
import time
import tkinter as tk
import unittest
import weakref
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

import excel_table_gui as gui


class TableWindowTests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.on_close = Mock()
        self.app = gui.TableToolsWindow(self.root, on_close=self.on_close)
        self.app.window.attributes("-alpha", 0)
        self.root.update()
        self.directory = tempfile.TemporaryDirectory(prefix="excel-table-gui-")
        self.folder = Path(self.directory.name)

    def tearDown(self):
        # Free every Tk object on its creator thread before a later worker GC.
        self.assertIs(threading.current_thread(), threading.main_thread())
        if self.app.running:
            self.app.request_cancel()
            self.finish()
        app_ref, root_ref = weakref.ref(self.app), weakref.ref(self.root)
        try:
            self.app.close()
            self.root.destroy()
        finally:
            self.app = None
            self.root = None
            gc.collect()
            self.directory.cleanup()
        self.assertIsNone(app_ref())
        self.assertIsNone(root_ref())

    def wait_for(self, condition):
        deadline = time.monotonic() + 10
        while not condition() and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertTrue(condition(), "Toplevel did not reach the expected state")

    def finish(self):
        self.wait_for(lambda: not self.app.running)

    def disabled(self, widget):
        return str(widget["state"]) == "disabled"

    def workbook(self, name="first.xlsx", headers=("分类", "数量"), rows=(("甲", 1), ("乙", 2))):
        path = self.folder / name
        book = Workbook()
        book.active.title = "数据表"
        book.active.append(list(headers))
        for row in rows:
            book.active.append(list(row))
        book.create_sheet("另一个表")
        book.save(path)
        book.close()
        return path

    def add(self, path):
        self.app.load_workbook(path)
        self.finish()
        self.app.sheet.set("数据表")
        self.app.add_source()

    def ready(self):
        self.add(self.workbook())
        self.app.preview()
        self.finish()
        self.assertFalse(self.disabled(self.app.run_button))

    def fake_result(self):
        directory = self.folder / "结果"
        return gui.tables.TaskResult(directory, (directory / "合表.xlsx",), 2, directory / "manifest.json")

    def test_explicit_sheet_header_and_inspection_off_main_thread(self):
        calls = []
        path = self.workbook()

        def inspect(source):
            calls.append((threading.current_thread(), source))
            return ["数据表", "另一个表"]

        with patch.object(gui.tables, "inspect_workbook", side_effect=inspect):
            self.app.load_workbook(path)
            self.finish()
        self.assertIsNot(calls[0][0], threading.main_thread())
        self.assertEqual(calls[0][1], path)
        self.assertEqual(self.app.sheet.get(), "")
        self.assertTrue(self.disabled(self.app.add_button))
        self.app.add_source()
        self.assertEqual(self.app.sources, [])
        self.app.sheet.set("另一个表")
        self.app.header_row.set("0")
        self.app.add_source()
        self.assertEqual(self.app.sources, [])
        self.app.header_row.set("3")
        self.app.add_source()
        self.assertEqual(self.app.sources, [gui.tables.TableSource(path, "另一个表", 3)])
        self.assertEqual(self.app.output_dir.get(), str(self.folder))
        self.app.add_source()
        self.assertEqual(len(self.app.sources), 1)
        self.assertIn("已在来源列表", self.app.status.get())

    def test_sources_invalidate_preview_but_output_and_mode_do_not(self):
        self.ready()
        preview = self.app.preview_result
        self.app.output_dir.set(str(self.folder / "other"))
        self.app.mode.set("split_rows")
        self.app._refresh_controls()
        self.assertIs(self.app.preview_result, preview)
        self.assertFalse(self.disabled(self.app.run_button))
        self.add(self.workbook("second.xlsx"))
        self.assertIsNone(self.app.preview_result)
        self.assertTrue(self.disabled(self.app.run_button))
        self.app.preview()
        self.finish()
        self.assertTrue(self.disabled(self.app.run_button), "Split must reject multiple sources")
        self.app.start()
        self.assertIn("一个来源", self.app.status.get())
        self.app.source_list.selection_set("1")
        self.app.remove_sources()
        self.assertIsNone(self.app.preview_result)
        self.assertEqual(len(self.app.sources), 1)

    def test_split_fields_and_row_limits_require_explicit_valid_parameters(self):
        self.ready()
        self.app.mode.set("split_column")
        self.app._refresh_controls()
        self.assertEqual(self.app.split_column.get(), "")
        self.assertEqual(tuple(self.app.column_combo["values"]), ("分类", "数量"))
        self.assertTrue(self.disabled(self.app.run_button))
        self.app.split_column.set("分类")
        self.app._refresh_controls()
        self.assertFalse(self.disabled(self.app.run_button))
        self.app.mode.set("split_rows")
        for invalid in ("", "abc", "0", "1048576", "1.5"):
            self.app.rows_per_file.set(invalid)
            self.app._refresh_controls()
            self.assertTrue(self.disabled(self.app.run_button), invalid)
        self.app.rows_per_file.set("1")
        self.app._refresh_controls()
        self.assertFalse(self.disabled(self.app.run_button))

    def test_output_default_follows_first_source_and_custom_choice_is_preserved(self):
        self.ready()
        self.app.source_list.selection_set("0")
        self.app.remove_sources()
        self.assertEqual(self.app.output_dir.get(), "")
        self.add(self.folder / "first.xlsx")
        self.assertEqual(self.app.output_dir.get(), str(self.folder))
        custom = self.folder / "custom"
        custom.mkdir()
        with patch.object(gui.filedialog, "askdirectory", return_value=str(custom)):
            self.app.choose_output()
        self.app.source_list.selection_set("0")
        self.app.remove_sources()
        self.add(self.folder / "first.xlsx")
        self.assertEqual(self.app.output_dir.get(), str(custom))

    def test_cancel_during_sheet_inspection_discards_the_returned_names(self):
        entered, release = threading.Event(), threading.Event()

        def inspect(path):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Test worker release timed out")
            return ["数据表"]

        with patch.object(gui.tables, "inspect_workbook", side_effect=inspect):
            self.app.load_workbook(self.folder / "first.xlsx")
            try:
                self.wait_for(entered.is_set)
                self.app.request_cancel()
            finally:
                release.set()
                self.finish()
        self.assertEqual(self.app._sheet_names, ())
        self.assertIsNone(self.app._inspected_path)
        self.assertTrue(self.disabled(self.app.add_button))
        self.assertIn("已取消", self.app.status.get())

    def test_frozen_snapshot_progress_and_controls_in_worker(self):
        self.ready()
        received = queue.Queue()
        release = threading.Event()
        result = self.fake_result()

        def execute(**kwargs):
            received.put((threading.current_thread(), kwargs))
            kwargs["progress"](dict(phase="writing", sheet="数据表", completed=4321, total=None, unit="rows"))
            if not release.wait(5):
                raise RuntimeError("Test worker release timed out")
            return result

        with patch.object(gui.tables, "run_tables", side_effect=execute):
            self.app.start()
            try:
                self.wait_for(lambda: "4,321 行" in self.app.status.get())
                thread, snapshot = received.get_nowait()
                self.assertIsNot(thread, threading.main_thread())
                self.assertEqual(snapshot["sources"], tuple(self.app.sources))
                self.assertEqual(snapshot["output_dir"], self.folder)
                self.assertTrue(snapshot["provenance"])
                self.assertNotIn("%", self.app.status.get())
                self.app.events.put((self.app.job, "progress", dict(phase="filling", completed=4322, total=None, unit="rows")))
                self.wait_for(lambda: "4,322 行" in self.app.status.get())
                self.assertIn("写入", self.app.status.get())
                self.assertTrue(self.disabled(self.app.choose_button))
                self.assertTrue(self.disabled(self.app.output_button))
                self.assertTrue(self.disabled(self.app.run_button))
                self.assertTrue(self.disabled(self.app.preview_button))
                self.assertTrue(self.disabled(self.app.provenance_check))
                self.app.provenance.set(False)
                self.app.output_dir.set(str(self.folder / "future"))
                self.assertTrue(snapshot["provenance"])
                self.assertEqual(snapshot["output_dir"], self.folder)
                self.app.started_at -= 65
                self.wait_for(lambda: "已用时 01:" in self.app.status.get())
            finally:
                release.set()
                self.finish()
        self.assertIs(self.app.result, result)
        self.assertIn("处理完成", self.app.status.get())
        self.assertFalse(self.disabled(self.app.open_button))
        terminal = self.app.status.get()
        self.app.events.put((self.app.job, "progress", {"completed": 999999}))
        self.root.after(150, lambda: None)
        self.wait_for(lambda: self.app.events.empty())
        self.assertEqual(self.app.status.get(), terminal)

    def test_busy_close_requests_cancel_keeps_both_windows_and_recovers(self):
        self.ready()
        entered = threading.Event()
        release = threading.Event()

        def execute(**kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Test worker release timed out")
            if kwargs["cancel"].is_set():
                raise gui.tables.TableCancelled("已取消，没有正式结果。")
            self.fail("Cancellation was not delivered")

        with patch.object(gui.tables, "run_tables", side_effect=execute):
            self.app.start()
            try:
                self.wait_for(entered.is_set)
                self.app.close()
                self.assertTrue(self.app.running)
                self.assertTrue(self.app.cancel_event.is_set())
                self.assertEqual(self.app.window.winfo_exists(), 1)
                self.assertEqual(self.root.winfo_exists(), 1)
                self.on_close.assert_not_called()
                self.assertIn("等待后台结束", self.app.status.get())
                self.app.events.put((self.app.job, "progress", dict(phase="saving", completed=1)))
                self.wait_for(lambda: self.app.events.empty())
                self.assertIn("已请求停止", self.app.status.get())
            finally:
                release.set()
                self.finish()
        self.assertIsNone(self.app.result)
        self.assertIn("已取消", self.app.status.get())
        self.assertFalse(self.disabled(self.app.run_button))
        self.app.close()
        self.app.close()
        self.on_close.assert_called_once_with()
        self.assertEqual(self.root.winfo_exists(), 1)

    def test_late_cancel_preserves_published_result_and_real_error_is_not_cancel(self):
        self.ready()
        result = self.fake_result()

        def committed(**kwargs):
            kwargs["cancel"].set()
            return result

        with patch.object(gui.tables, "run_tables", side_effect=committed):
            self.app.start()
            self.finish()
        self.assertIs(self.app.result, result)
        self.assertIn("处理完成", self.app.status.get())

        def failed(**kwargs):
            kwargs["cancel"].set()
            error = OSError("synthetic write failed")
            error.add_note("临时文件清理失败：synthetic.part")
            raise error

        with patch.object(gui.tables, "run_tables", side_effect=failed):
            self.app.start()
            self.finish()
        self.assertIn("处理失败", self.app.status.get())
        self.assertIs(self.app.result, result, "A failed new task must not lose the previous successful directory")
        self.assertIn("临时文件清理失败", self.app.log.get("1.0", "end"))
        with patch.object(gui.os, "startfile", create=True) as open_folder:
            self.app.open_result()
            open_folder.assert_called_once_with(str(result.directory))

    def test_preview_and_inspection_failures_leave_window_usable(self):
        self.ready()
        with patch.object(gui.tables, "preview_sources", side_effect=ValueError("表头不一致")):
            self.app.preview()
            self.finish()
        self.assertIsNone(self.app.preview_result)
        self.assertTrue(self.disabled(self.app.run_button))
        self.assertFalse(self.disabled(self.app.preview_button))
        self.assertIn("表头不一致", self.app.log.get("1.0", "end"))
        with patch.object(gui.tables, "inspect_workbook", side_effect=PermissionError("locked")):
            self.app.load_workbook(self.folder / "locked.xlsx")
            self.finish()
        self.assertEqual(self.app.sheet.get(), "")
        self.assertTrue(self.disabled(self.app.add_button))
        self.assertFalse(self.disabled(self.app.choose_button))
        self.assertIn("检查读写权限", self.app.log.get("1.0", "end"))

    def test_thread_creation_and_start_failure_restore_idle(self):
        self.ready()
        for phase in ("construct", "start"):
            for error_type in (MemoryError, RuntimeError):
                with self.subTest(phase=phase, error=error_type.__name__):
                    thread = Mock()
                    target = error_type("synthetic thread failure")
                    if phase == "start":
                        thread.start.side_effect = target
                    with patch.object(gui.threading, "Thread", side_effect=target if phase == "construct" else None, return_value=thread):
                        self.app.start()
                    self.assertFalse(self.app.running)
                    self.assertIsNone(self.app.result)
                    self.assertFalse(self.disabled(self.app.run_button))
                    self.assertTrue(self.disabled(self.app.stop_button))
        with patch.object(gui.tables, "run_tables", return_value=self.fake_result()):
            self.app.start()
            self.finish()
        self.assertIn("处理完成", self.app.status.get())

    def test_minimum_window_keeps_actions_and_status_visible(self):
        self.app.window.geometry("760x700")
        self.root.update()
        width, height = self.app.window.winfo_width(), self.app.window.winfo_height()
        for widget in (self.app.choose_button, self.app.add_button, self.app.remove_button,
                       self.app.preview_button, self.app.run_button, self.app.stop_button,
                       self.app.open_button, self.app.status_label):
            x = widget.winfo_rootx() - self.app.window.winfo_rootx()
            y = widget.winfo_rooty() - self.app.window.winfo_rooty()
            self.assertGreaterEqual(x, 0)
            self.assertGreaterEqual(y, 0)
            self.assertLessEqual(x + widget.winfo_width(), width)
            self.assertLessEqual(y + widget.winfo_height(), height)

    def test_real_merge_then_split_rows_and_group_results(self):
        first, second = self.workbook(), self.workbook("second.xlsx", rows=(("甲", 3),))
        self.add(first)
        self.add(second)
        self.app.preview()
        self.finish()
        self.app.start()
        self.finish()
        self.assertIsNotNone(self.app.result, self.app.log.get("1.0", "end"))
        self.assertEqual(self.app.result.rows, 3)
        book = load_workbook(self.app.result.files[0], read_only=True)
        try:
            rows = list(book.active.values)
            self.assertEqual([row[:2] for row in rows], [("分类", "数量"), ("甲", 1), ("乙", 2), ("甲", 3)])
            self.assertEqual(rows[0][-3:], ("来源文件", "来源工作表", "来源行号"))
        finally:
            book.close()
        self.app.source_list.selection_set("1")
        self.app.remove_sources()
        self.app.preview()
        self.finish()
        self.app.mode.set("split_rows")
        self.app.rows_per_file.set("1")
        self.app.provenance.set(False)
        self.app.start()
        self.finish()
        self.assertEqual(len(self.app.result.files), 2)
        self.assertEqual(self.app.result.rows, 2)
        self.app.mode.set("split_column")
        self.app.split_column.set("分类")
        self.app.start()
        self.finish()
        self.assertEqual(len(self.app.result.files), 2)
        self.assertEqual(self.app.result.rows, 2)
        self.assertTrue(self.app.result.manifest.exists())

    def test_formula_error_preserves_input_has_no_output_and_can_retry(self):
        path = self.workbook(rows=(("甲", "=1+1"),))
        original = path.read_bytes()
        self.add(path)
        self.app.preview()
        self.finish()
        self.app.start()
        self.finish()
        self.assertIn("处理失败", self.app.status.get())
        self.assertIn("公式", self.app.log.get("1.0", "end"))
        self.assertIsNone(self.app.result)
        self.assertEqual(list(self.folder.iterdir()), [path])
        self.assertEqual(path.read_bytes(), original)
        self.workbook(rows=(("甲", 2),))
        self.app.start()
        self.finish()
        self.assertIsNotNone(self.app.result, self.app.log.get("1.0", "end"))
        self.assertEqual(self.app.result.rows, 1)


if __name__ == "__main__":
    unittest.main()
