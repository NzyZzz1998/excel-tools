"""Independent scalar XLSX fixtures for the new table reader and writer."""
import errno
from itertools import count
from pathlib import Path
import tempfile
import threading
import unittest
from unittest import mock
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import excel_table_io as io


NS = io.NS


def text_cell(ref, value, style=None):
    attr = '' if style is None else f' s="{style}"'
    content = escape(value).replace('\r', '&#13;')
    return f'<c r="{ref}" t="inlineStr"{attr}><is><t xml:space="preserve">{content}</t></is></c>'


def value_cell(ref, value, kind='n', style=None, extra=''):
    attr = '' if style is None else f' s="{style}"'
    payload = '' if value is None else '<v>' + escape(value) + '</v>'
    return f'<c r="{ref}" t="{kind}"{attr}>{extra}{payload}</c>'


def row(number, *cells, attributes=''):
    return f'<row r="{number}"{attributes}>' + ''.join(cells) + '</row>'


def makebook(path, rows, *, sst=None, merges=(), date1904=False, strict=False,
             styles=None, before='', after=''):
    ns = 'http://purl.oclc.org/ooxml/spreadsheetml/main' if strict else NS
    rel = io.REL
    types = ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
             '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
             '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
             + ('<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>' if sst is not None else '') + '</Types>')
    if styles is None:
        styles = (f'<styleSheet xmlns="{ns}"><numFmts count="1"><numFmt numFmtId="164" formatCode="000000"/></numFmts>'
                  '<fonts count="1"><font/></fonts><fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border/></borders>'
                  '<cellStyleXfs count="1"><xf numFmtId="0"/></cellStyleXfs>'
                  '<cellXfs count="3"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="164"/></cellXfs>'
                  '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>')
    with ZipFile(path, 'x', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', types)
        archive.writestr('_rels/.rels', f'<Relationships xmlns="{io.PKG_REL}"><Relationship Id="r1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        archive.writestr('xl/workbook.xml', f'<workbook xmlns="{ns}" xmlns:r="{rel}"><workbookPr date1904="{int(date1904)}"/><sheets><sheet name="synthetic" sheetId="1" r:id="r1"/></sheets></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels', f'<Relationships xmlns="{io.PKG_REL}"><Relationship Id="r1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="r2" Type="{rel}/styles" Target="styles.xml"/>'
                         + (f'<Relationship Id="r3" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/>' if sst is not None else '') + '</Relationships>')
        archive.writestr('xl/styles.xml', styles)
        if sst is not None:
            archive.writestr('xl/sharedStrings.xml', f'<sst xmlns="{ns}">' + sst + '</sst>')
        merged = '<mergeCells>' + ''.join(f'<mergeCell ref="{ref}"/>' for ref in merges) + '</mergeCells>' if merges else ''
        archive.writestr('xl/worksheets/sheet1.xml', f'<worksheet xmlns="{ns}"><dimension ref="A1"/>' + before + '<sheetData>' + ''.join(rows) + '</sheetData>' + merged + after + '</worksheet>')
    return path


class TableIOTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='excel-table-test-')
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_scalar_roundtrip_retains_raw_numbers_formats_text_and_epoch(self):
        values = [io.Cell('n', '123456789012345678901'), io.Cell('s', '00123'),
                  io.Cell('n', '12', '000000'), io.Cell('n', '45000.123456789', 14),
                  io.Cell('s', '=SUM(A1:A2)'), io.Cell('s', '#N/A'), io.Cell('e', '#N/A'),
                  io.Cell('b', '0'), io.Cell('d', '2026-10-11T12:34:56'), io.Cell('s', ''),
                  io.Cell('s', ' line\r\n\t<&> _x0041_ \x01 \U0001f600 ')]
        target = self.root/'types.xlsx'
        with io.XlsxWriter(target, [f'col{i}' for i in range(len(values))], date1904=True) as writer:
            writer.append(values)
        with io.TableReader(target, '数据') as reader:
            self.assertTrue(reader.date1904)
            self.assertEqual(list(reader.iter_rows()), [(2, values)])
        with ZipFile(target) as archive:
            self.assertIsNone(archive.testzip())
            raw = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(raw.find(f'{{{NS}}}dimension').get('ref'), 'A1:K2')
            self.assertFalse(list(raw.iter(f'{{{NS}}}f')))
        wb = openpyxl.load_workbook(target, read_only=True)
        try:
            self.assertEqual(wb['数据']['E2'].value, '=SUM(A1:A2)')
            self.assertEqual(wb['数据']['E2'].data_type, 's')
            self.assertEqual(wb['数据']['C2'].number_format, '000000')
        finally:
            wb.close()

    def test_wrong_dimension_sparse_rows_and_explicit_blank_values(self):
        path = makebook(self.root/'sparse.xlsx', [
            row(1, text_cell('A1', ' one '), text_cell('B1', 'two')),
            row(2, value_cell('A2', None)), row(4, text_cell('B4', '')),
            row(6, value_cell('A6', '0'), value_cell('B6', '0', 'b')),
            row(9, text_cell('A9', ' '))])
        with io.TableReader(path, 'synthetic') as reader:
            self.assertEqual(reader.headers, ['one', 'two'])
            rows = list(reader.iter_rows())
        self.assertEqual([number for number, _ in rows], [4, 6, 9])
        self.assertEqual(rows[0][1], [None, io.Cell('s', '')])
        self.assertEqual(rows[-1][1], [io.Cell('s', ' '), None])

    def test_shared_strings_rich_text_escapes_and_bounded_cache(self):
        entries = ['<si><r><t> head</t></r><r><t>er </t></r><rPh sb="0" eb="1"><t>ignore</t></rPh></si>']
        entries += ['<si><t>' + str(i) + '</t></si>' for i in range(1, 1101)]
        entries.append('<si><t>_x0041_ _x005F_x0041_ _xD83D__xDE00_</t></si>')
        path = makebook(self.root/'shared.xlsx', [row(i + 1, value_cell(f'A{i+1}', str(i), 's')) for i in range(len(entries))], sst=''.join(entries))
        with io.TableReader(path, 'synthetic') as reader:
            self.assertEqual(reader.headers, ['header'])
            rows = list(reader.iter_rows())
            self.assertEqual(rows[-1][1][0].value, 'A _x0041_ \U0001f600')
            self.assertLessEqual(len(reader._strings.cache), 1024)
            self.assertLessEqual(reader._strings.cache_bytes, 4 * 1024 * 1024)
            self.assertEqual(reader._strings.count, len(entries))

    def test_formula_with_or_without_cache_is_rejected_in_header_or_data(self):
        for number, cached in ((1, '1'), (1, None), (2, '1'), (2, None)):
            with self.subTest(number=number, cached=cached):
                rows = [row(1, text_cell('A1', 'header'))] if number == 2 else []
                rows.append(row(number, value_cell(f'A{number}', cached, extra='<f>1+0</f>')))
                path = makebook(self.root/f'formula-{number}-{cached}.xlsx', rows)
                with self.assertRaisesRegex(ValueError, f'{path.name} / synthetic / A{number}.*公式'):
                    with io.TableReader(path, 'synthetic') as reader:
                        list(reader.iter_rows())

    def test_merge_rejected_on_exhaustion_but_only_when_reaching_selected_range(self):
        rows = [row(1, text_cell('A1', 'title')), row(3, text_cell('A3', 'header')), row(4, text_cell('A4', 'value'))]
        for ref, rejected in (('A1:C1', False), ('A1:A3', True), ('A4:A5', True)):
            with self.subTest(ref=ref):
                path = makebook(self.root/('merge-' + ref.replace(':', '-') + '.xlsx'), rows, merges=[ref])
                with io.TableReader(path, 'synthetic', header_row=3) as reader:
                    iterator = reader.iter_rows()
                    self.assertEqual(next(iterator)[0], 4)
                    if rejected:
                        with self.assertRaisesRegex(ValueError, '拆分填充'):
                            next(iterator)
                    else:
                        self.assertEqual(list(iterator), [])

    def test_header_duplicates_blanks_missing_and_numeric_header(self):
        for name, cells in [('duplicate', [text_cell('A1', 'a'), text_cell('B1', ' a ')]),
                            ('empty', [text_cell('A1', 'a'), text_cell('C1', 'c')]),
                            ('blank', [text_cell('A1', ' ')])]:
            path = makebook(self.root/(name + '.xlsx'), [row(1, *cells)])
            with self.subTest(name=name), self.assertRaises(ValueError):
                with io.TableReader(path, 'synthetic'):
                    pass
        path = makebook(self.root/'numeric-header.xlsx', [row(1, value_cell('A1', '123'))])
        with io.TableReader(path, 'synthetic') as reader:
            self.assertEqual(reader.headers, ['123'])
            self.assertEqual(list(reader.iter_rows()), [])
        with self.assertRaisesRegex(ValueError, '表头行'):
            with io.TableReader(path, 'synthetic', header_row=2):
                pass

    def test_extra_nonblank_column_rejected_but_empty_styled_cell_allowed(self):
        for value in (None, '0', ''):
            path = makebook(self.root/f'extra-{value!s}.xlsx', [row(1, text_cell('A1', 'a')),
                row(2, text_cell('A2', 'v'), value_cell('C2', value, 'str' if value == '' else 'n', style=1))])
            with self.subTest(value=value), io.TableReader(path, 'synthetic') as reader:
                if value is None:
                    self.assertEqual(len(list(reader.iter_rows())), 1)
                else:
                    with self.assertRaisesRegex(ValueError, '超出表头列数'):
                        list(reader.iter_rows())

    def test_cell_metadata_rejected_even_without_value(self):
        variants = ['<c r="A2" vm="1"/>', '<c r="A2" cm="1"/>', '<c r="A2"><extLst/></c>']
        for i, cell in enumerate(variants):
            path = makebook(self.root/f'metadata-{i}.xlsx', [row(1, text_cell('A1', 'a')), row(2, cell)])
            with io.TableReader(path, 'synthetic') as reader, self.assertRaisesRegex(ValueError, '元数据'):
                list(reader.iter_rows())

    def test_strict_namespace_and_inherited_number_formats(self):
        path = makebook(self.root/'strict.xlsx', [row(1, text_cell('A1', 'a'), text_cell('B1', 'b')),
            row(2, value_cell('A2', '12'), value_cell('B2', '44000', style=1)),
            row(3, value_cell('A3', '44001'), attributes=' s="1" customFormat="1"')],
            before='<cols><col min="1" max="1" style="2"/></cols>', strict=True)
        self.assertEqual(io.list_sheets(path), ['synthetic'])
        with io.TableReader(path, 'synthetic') as reader:
            rows = list(reader.iter_rows())
        self.assertEqual([cell.number_format for cell in rows[0][1]], ['000000', 14])
        self.assertEqual(rows[1][1][0].number_format, 14)

    def test_date_format_recognizer_ignores_literals_and_decorations(self):
        for value in (14, 22, 45, 58, 'yyyy-mm-dd', 'h:mm', '[h]:mm:ss', '[$-409]mmm d, yyyy'):
            self.assertTrue(io.is_date_format(value), value)
        for value in (0, 1, 164, '000000', '0.00', '0 "days"', '0\\d', '0_m', '[Red]0', '[>=1]0'):
            self.assertFalse(io.is_date_format(value), value)

    def test_cancel_during_shared_strings_rows_and_zip_save_releases_resources(self):
        event = threading.Event()
        event.set()
        path = makebook(self.root/'cancel.xlsx', [row(1, text_cell('A1', 'a')), row(2, text_cell('A2', 'b'))])
        with self.assertRaises(io.TableCancelled):
            with io.TableReader(path, 'synthetic', cancel=event):
                pass
        event.clear()
        with io.TableReader(path, 'synthetic', cancel=event) as reader:
            event.set()
            with self.assertRaises(io.TableCancelled):
                list(reader.iter_rows())
        event.clear()
        output = self.root/'cancel-output.xlsx'
        def cancel_on_save(detail):
            if detail['phase'] == 'saving' and detail['completed']:
                event.set()
        with self.assertRaises(io.TableCancelled):
            with io.XlsxWriter(output, ['a'], cancel=event, progress=cancel_on_save) as writer:
                writer.append([io.Cell('s', 'v')])
        self.assertFalse(output.exists())
        path.unlink()  # Windows detects leaked archive handles here.

    def test_cancel_while_loading_shared_strings_removes_disk_index(self):
        path = makebook(self.root/'cancel-sst.xlsx', [row(1, value_cell('A1', '0', 's'))],
                        sst='<si><t>head</t></si>' * 20)
        event = threading.Event()
        calls, directories = [], []
        def progress(detail):
            calls.append(detail)
            if len(calls) == 2:
                event.set()
        original_temp = io.TemporaryDirectory
        def capture_temp(*args, **kwargs):
            directory = original_temp(*args, **kwargs)
            directories.append(Path(directory.name))
            return directory
        with mock.patch.object(io, 'TemporaryDirectory', side_effect=capture_temp), mock.patch.object(io.time, 'monotonic', side_effect=count()), self.assertRaises(io.TableCancelled):
            with io.TableReader(path, 'synthetic', progress=progress, cancel=event):
                pass
        self.assertEqual(len(calls), 2)
        self.assertTrue(directories)
        self.assertTrue(all(not directory.exists() for directory in directories))
        path.unlink()

    def test_invalid_scalar_payloads_rejected_with_coordinates(self):
        for i, (kind, value) in enumerate((('n', 'NaN'), ('n', '1e309'), ('n', '1_000'),
                                          ('b', '2'), ('d', '2026-02-30'), ('d', 'not-a-date'))):
            path = makebook(self.root/f'invalid-{i}.xlsx', [row(1, text_cell('A1', 'a')),
                            row(2, value_cell('A2', value, kind))])
            with self.subTest(kind=kind, value=value), io.TableReader(path, 'synthetic') as reader:
                with self.assertRaisesRegex(ValueError, 'synthetic / A2'):
                    list(reader.iter_rows())
            with self.assertRaisesRegex(ValueError, 'A2'):
                with io.XlsxWriter(self.root/f'invalid-out-{i}.xlsx', ['a']) as writer:
                    writer.append([io.Cell(kind, value)])

    def test_large_layout_tail_is_released_while_metadata_inside_cell_stays_visible(self):
        tail = '<hyperlinks>' + '<hyperlink ref="A2" location="A1"/>' * 10000 + '</hyperlinks>'
        path = makebook(self.root/'tail.xlsx', [row(1, text_cell('A1', 'a')), row(2, text_cell('A2', 'v'))], after=tail)
        original = io.ET.iterparse
        largest = [0]
        def observed(*args, **kwargs):
            for event, element in original(*args, **kwargs):
                yield event, element
                if event == 'end' and io._local(element.tag) == 'hyperlinks':
                    largest[0] = len(element)
        with mock.patch.object(io.ET, 'iterparse', observed), io.TableReader(path, 'synthetic') as reader:
            self.assertEqual(len(list(reader.iter_rows())), 1)
        self.assertEqual(largest[0], 0)

    def test_writer_limits_exclusive_creation_and_body_error_cleanup(self):
        path = self.root/'existing.xlsx'
        path.write_bytes(b'old')
        with self.assertRaises(FileExistsError):
            with io.XlsxWriter(path, ['a']) as writer:
                writer.append([io.Cell('s', 'new')])
        self.assertEqual(path.read_bytes(), b'old')
        target = self.root/'too-many.xlsx'
        with mock.patch.object(io, 'MAX_DATA_ROWS', 1), self.assertRaisesRegex(ValueError, '最多容纳'):
            with io.XlsxWriter(target, ['a']) as writer:
                writer.append([io.Cell('n', '1')])
                writer.append([io.Cell('n', '2')])
        self.assertFalse(target.exists())
        self.assertTrue(writer._rows.closed)

    def test_real_partial_zip_write_failure_retains_original_error_and_cleans_output(self):
        path = self.root/'write-failed.xlsx'
        error = OSError(errno.ENOSPC, 'synthetic full disk')
        original = io.ZipFile.open
        written = []
        def injected(archive, name, mode='r', *args, **kwargs):
            stream = original(archive, name, mode, *args, **kwargs)
            if mode == 'w' and name == 'xl/worksheets/sheet1.xml':
                original_write = stream.write
                def fail(data):
                    written.append(original_write(data[:5]))
                    raise error
                stream.write = fail
            return stream
        with mock.patch.object(io.ZipFile, 'open', injected), self.assertRaises(OSError) as raised:
            with io.XlsxWriter(path, ['a']) as writer:
                writer.append([io.Cell('s', 'synthetic')])
        self.assertIs(raised.exception, error)
        self.assertEqual(sum(written), 5)
        self.assertFalse(path.exists())
        self.assertTrue(writer._rows.closed)

    def test_progress_is_bounded_and_observer_failure_does_not_change_output(self):
        events = []
        path = self.root/'progress.xlsx'
        with mock.patch.object(io.time, 'monotonic', return_value=0):
            with io.XlsxWriter(path, ['a'], progress=events.append) as writer:
                for i in range(500):
                    writer.append([io.Cell('n', str(i))])
            with io.TableReader(path, '数据', progress=events.append) as reader:
                self.assertEqual(len(list(reader.iter_rows())), 500)
        self.assertLessEqual(len(events), 6)
        def broken(_):
            raise ValueError('synthetic observer failure')
        with io.TableReader(path, '数据', progress=broken) as reader:
            self.assertEqual(len(list(reader.iter_rows())), 500)


if __name__ == '__main__':
    unittest.main()
