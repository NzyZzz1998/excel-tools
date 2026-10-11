"""Independent v1.2 table-operation acceptance tests using openpyxl as oracle.

Fixtures are synthetic. Expected row mappings are specified here rather than
derived through application readers/grouping helpers. This module owns no real
business input and does not exercise the legacy unmerge GUI.
"""
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import errno
import hashlib
import json
import os
import sqlite3
from pathlib import Path
import tempfile
import threading
import unittest
from unittest import mock
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904

import excel_table_tools as mod


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def edit_sheet(path, edit):
    """Patch only a small synthetic source to express cached formulas/dimensions."""
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    sheet = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    edit(sheet)
    members["xl/worksheets/sheet1.xml"] = ET.tostring(sheet, encoding="utf-8")
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, contents in members.items():
            archive.writestr(name, contents)


class TableToolsAcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="excel-table-acceptance-")
        self.root = Path(self.temporary.name)
        self.inputs = self.root / "inputs"
        self.output = self.root / "results"
        self.inputs.mkdir()
        self.output.mkdir()

    def tearDown(self):
        self.temporary.cleanup()

    def book(self, name, headers, rows=(), *, sheet="数据", header_row=1,
             epoch=None, merges=()):
        path = self.inputs / name
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = sheet
        if epoch is not None:
            workbook.epoch = epoch
        if header_row > 1:
            worksheet.cell(1, 1, "SYNTHETIC preamble")
        for row_number, values in enumerate([headers, *rows], header_row):
            for column_number, value in enumerate(values, 1):
                cell = worksheet.cell(row_number, column_number, value)
                # All ordinary fixture strings are literal text, including '='.
                if isinstance(value, str):
                    cell.data_type = "s"
        for ref in merges:
            worksheet.merge_cells(ref)
        workbook.save(path)
        workbook.close()
        return path

    def source(self, path, sheet="数据", header_row=1):
        return mod.TableSource(path, sheet, header_row=header_row)

    def run_job(self, sources, **options):
        options.setdefault("provenance", False)
        return mod.run_tables(sources, self.output, **options)

    def read_output(self, path):
        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip(), str(path))
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertEqual(len(workbook.worksheets), 1)
            worksheet = workbook.worksheets[0]
            rows = [tuple((cell.value, cell.data_type, cell.number_format)
                          for cell in row) for row in worksheet.iter_rows()]
        finally:
            workbook.close()
        return rows

    def value_rows(self, result):
        rows = []
        for path in result.files:
            cells = self.read_output(path)
            rows.extend(tuple(cell[0] for cell in row) for row in cells[1:])
        return rows

    def assert_manifest(self, result, expected_rows):
        self.assertIsInstance(result.directory, Path)
        self.assertTrue(result.directory.is_absolute())
        # A uniquely reserved task parent contains one atomically published 结果.
        self.assertEqual(result.directory.parent.parent.resolve(), self.output.resolve())
        self.assertEqual(result.directory.name, "结果")
        self.assertIsInstance(result.files, tuple)
        self.assertIsInstance(result.manifest, Path)
        self.assertEqual(result.manifest.parent, result.directory)
        self.assertEqual(result.rows, expected_rows)
        data = json.loads(result.manifest.read_text(encoding="utf-8"))
        self.assertEqual(data["format_version"], 1)
        self.assertEqual(data["rows"], expected_rows)
        self.assertEqual(sum(entry["rows"] for entry in data["files"]), expected_rows)
        self.assertEqual({entry["name"] for entry in data["files"]},
                         {path.name for path in result.files})
        for path in result.files:
            self.assertEqual(path.parent, result.directory)
            self.assertEqual(path.suffix, ".xlsx")
        self.assertEqual(set(p.name for p in result.directory.iterdir()),
                         {p.name for p in result.files} | {result.manifest.name})
        return data

    def assert_rejected_cleanly(self, sources, **options):
        before = {p.relative_to(self.output): digest(p) if p.is_file() else None
                  for p in self.output.rglob("*")}
        with self.assertRaises(ValueError):
            self.run_job(sources, **options)
        after = {p.relative_to(self.output): digest(p) if p.is_file() else None
                 for p in self.output.rglob("*")}
        self.assertEqual(after, before)

    def test_preview_trims_header_and_merge_reorders_by_name(self):
        first = self.book("first.xlsx", [" 编号 ", "日期", "金额"],
                          [["0012", "2024-01-01", 0]])
        second = self.book("second.xlsx", ["金额", "编号", "日期"],
                           [[2.5, "0034", "2024-01-02"]])
        sources = [self.source(first), self.source(second)]
        hashes = [digest(first), digest(second)]
        preview = mod.preview_sources(sources)
        self.assertEqual(preview["headers"], ["编号", "日期", "金额"])
        self.assertEqual(len(preview["sources"]), 2)
        self.assertFalse(preview["date1904"])
        self.assertEqual(list(self.output.iterdir()), [])
        result = self.run_job(sources, mode="merge")
        self.assert_manifest(result, 2)
        self.assertEqual(self.value_rows(result),
                         [("0012", "2024-01-01", 0), ("0034", "2024-01-02", 2.5)])
        self.assertEqual([digest(first), digest(second)], hashes)

    def test_schema_mismatch_and_ambiguous_headers_do_not_publish(self):
        good = self.book("good.xlsx", ["A", "B"], [[1, 2]])
        for index, headers in enumerate((["A"], ["A", "B", "C"],
                                         ["A", " A "], ["A", None, "B"])):
            with self.subTest(headers=headers):
                bad = self.book("bad-{}.xlsx".format(index), headers, [[3] * len(headers)])
                self.assert_rejected_cleanly([self.source(good), self.source(bad)], mode="merge")

    def test_duplicate_source_is_rejected_but_different_sheets_are_distinct(self):
        path = self.book("multiple.xlsx", ["编号"], [["one"]])
        workbook = openpyxl.load_workbook(path)
        worksheet = workbook.create_sheet("其他")
        worksheet.append(["编号"])
        worksheet.append(["two"])
        workbook.save(path)
        workbook.close()
        self.assert_rejected_cleanly([self.source(path), self.source(path)], mode="merge")
        result = self.run_job([self.source(path), self.source(path, "其他")], mode="merge")
        self.assertEqual(self.value_rows(result), [("one",), ("two",)])

    def test_provenance_survives_same_filenames_blank_rows_and_name_collision(self):
        headers = ["来源文件", "编号"]
        first = self.book("a/same.xlsx", headers, [["original-a", "001"], [None, None],
                                                 ["original-b", "002"]], header_row=2)
        second = self.book("b/same.xlsx", headers, [["original-c", "003"]], header_row=2)
        result = self.run_job([self.source(first, header_row=2), self.source(second, header_row=2)],
                              mode="merge", provenance=True)
        manifest = self.assert_manifest(result, 3)
        self.assertEqual(manifest["provenance_columns"], ["来源文件_2", "来源工作表", "来源行号"])
        self.assertEqual(self.value_rows(result), [
            ("original-a", "001", str(first.resolve()), "数据", 3),
            ("original-b", "002", str(first.resolve()), "数据", 5),
            ("original-c", "003", str(second.resolve()), "数据", 3)])
        self.assertEqual({entry["path"] for entry in manifest["sources"]},
                         {str(first.resolve()), str(second.resolve())})

    def test_values_number_formats_and_literal_formula_text_survive(self):
        path = self.book("types.xlsx", ["文本编号", "显示编号", "日期", "布尔", "零", "文本", "错误"],
                         [["0012", 12, datetime(2024, 1, 2, 3, 4, 5), False, 0,
                           "=1+1", "#DIV/0!"],
                          ["123456789012345678901", None, None, True, 0.25,
                           "SYNTHETIC_CONTROL_TEXT", "#N/A"]])
        workbook = openpyxl.load_workbook(path)
        worksheet = workbook.active
        worksheet["B2"].number_format = "0000"
        worksheet["C2"].number_format = "yyyy-mm-dd hh:mm:ss"
        worksheet["G2"].data_type = "e"
        # G3 deliberately remains a literal string resembling an error.
        workbook.save(path)
        workbook.close()
        # Some openpyxl XML backends serialize literal CR; constructing &#13;
        # explicitly avoids losing the intended character in the fixture itself.
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
            b"SYNTHETIC_CONTROL_TEXT", b"  first&#13;second\nthird\t  ")
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            for name, contents in members.items():
                archive.writestr(name, contents)
        original = openpyxl.load_workbook(path, read_only=True)
        try:
            self.assertEqual(original.active["F3"].value, "  first\rsecond\nthird\t  ")
        finally:
            original.close()
        result = self.run_job([self.source(path)], mode="merge")
        cells = self.read_output(result.files[0])
        self.assertEqual(cells[1][0], ("0012", "s", "General"))
        self.assertEqual(cells[1][1], (12, "n", "0000"))
        self.assertEqual(cells[1][2], (datetime(2024, 1, 2, 3, 4, 5), "d", "yyyy-mm-dd hh:mm:ss"))
        self.assertEqual(cells[1][3][0:2], (False, "b"))
        self.assertEqual(cells[1][4][0:2], (0, "n"))
        self.assertEqual(cells[1][5][0:2], ("=1+1", "s"))
        self.assertEqual(cells[1][6][0:2], ("#DIV/0!", "e"))
        self.assertEqual(cells[2][0][0], "123456789012345678901")
        self.assertEqual(cells[2][5][0], "  first\rsecond\nthird\t  ")
        self.assertEqual(cells[2][6][0:2], ("#N/A", "s"))

    def test_formulas_with_or_without_cache_in_header_or_data_are_rejected(self):
        for coordinate in ("A1", "A2"):
            for cached in (False, True):
                with self.subTest(coordinate=coordinate, cached=cached):
                    path = self.book("formula-{}-{}.xlsx".format(coordinate, cached), ["ID"], [[1]])
                    def add_formula(sheet):
                        cell = sheet.find(".//{%s}c[@r='%s']" % (NS, coordinate))
                        cell.clear()
                        cell.set("r", coordinate)
                        ET.SubElement(cell, "{%s}f" % NS).text = "1+1"
                        if cached:
                            ET.SubElement(cell, "{%s}v" % NS).text = "2"
                    edit_sheet(path, add_formula)
                    self.assert_rejected_cleanly([self.source(path)], mode="merge")

    def test_merge_regions_above_header_are_ignored_but_data_or_header_merges_reject(self):
        allowed = self.book("title.xlsx", ["ID", "Value"], [[1, 2]], header_row=2,
                            merges=["A1:B1"])
        result = self.run_job([self.source(allowed, header_row=2)], mode="merge")
        self.assertEqual(self.value_rows(result), [(1, 2)])
        for name, region in (("header", "A2:B2"), ("data", "A3:A4")):
            with self.subTest(region=region):
                bad = self.book(name + ".xlsx", ["ID", "Value"], [[1, 2], [None, 3]],
                                header_row=2, merges=[region])
                self.assert_rejected_cleanly([self.source(allowed, header_row=2),
                                              self.source(bad, header_row=2)], mode="merge")

    def test_mixed_date_epochs_reject_but_single_1904_source_preserves_date(self):
        date = datetime(2024, 2, 29)
        first = self.book("windows.xlsx", ["Date"], [[date]])
        second = self.book("mac.xlsx", ["Date"], [[date]], epoch=CALENDAR_MAC_1904)
        self.assert_rejected_cleanly([self.source(first), self.source(second)], mode="merge")
        result = self.run_job([self.source(second)], mode="merge")
        self.assertEqual(self.value_rows(result), [(date,)])

    def test_inaccurate_dimension_does_not_truncate_real_rows_or_columns(self):
        path = self.book("wrong-dimension.xlsx", ["ID", "B", "C"], [[1, 2, 3], [4, 5, 6]])
        edit_sheet(path, lambda sheet: sheet.find("{%s}dimension" % NS).set("ref", "A1"))
        result = self.run_job([self.source(path)], mode="merge")
        self.assertEqual(self.value_rows(result), [(1, 2, 3), (4, 5, 6)])

    def test_blank_rows_skip_without_dropping_false_zero_or_whitespace(self):
        path = self.book("blank-rows.xlsx", ["Value"], [[None], [0], [False], [" "], [None]])
        result = self.run_job([self.source(path)], mode="merge")
        self.assert_manifest(result, 3)
        rows = self.read_output(result.files[0])[1:]
        self.assertEqual([(r[0][0], r[0][1]) for r in rows], [(0, "n"), (False, "b"), (" ", "s")])

    def test_header_only_sources_have_one_valid_empty_output_for_each_mode(self):
        path = self.book("empty.xlsx", ["ID", "Value"])
        for mode in ("merge", "split_rows", "split_column"):
            with self.subTest(mode=mode):
                result = self.run_job([self.source(path)], mode=mode, split_column="ID")
                self.assert_manifest(result, 0)
                self.assertEqual(len(result.files), 1)
                self.assertEqual(len(self.read_output(result.files[0])), 1)

    def test_split_rows_boundaries_repeat_header_once_and_conserve_order(self):
        for count in (2, 3, 4, 6, 7):
            with self.subTest(count=count):
                path = self.book("rows-{}.xlsx".format(count), ["ID"], [[i] for i in range(count)])
                result = self.run_job([self.source(path)], mode="split_rows", rows_per_file=3)
                self.assert_manifest(result, count)
                self.assertEqual(len(result.files), (count + 2) // 3)
                self.assertEqual(self.value_rows(result), [(i,) for i in range(count)])
                for output in result.files:
                    rows = self.read_output(output)
                    self.assertEqual(rows[0][0][0], "ID")
                    self.assertLessEqual(len(rows) - 1, 3)

    def test_merge_and_group_rollover_keep_every_row_once(self):
        path = self.book("rollover.xlsx", ["Group", "ID"], [["same", i] for i in range(7)])
        for mode in ("merge", "split_column"):
            with self.subTest(mode=mode), mock.patch.object(mod, "MAX_DATA_ROWS", 3):
                result = self.run_job([self.source(path)], mode=mode, split_column="Group")
                self.assert_manifest(result, 7)
                self.assertEqual(len(result.files), 3)
                self.assertEqual(self.value_rows(result), [("same", i) for i in range(7)])
                self.assertTrue(all(len(self.read_output(p)) - 1 <= 3 for p in result.files))

    def test_grouping_is_typed_numeric_canonical_and_independent_of_format(self):
        rows = [[1, "n-one"], [1.0, "n-one-formatted"], [True, "bool"],
                ["1", "text"], ["001", "identifier"], [None, "blank"], [" ", "space"]]
        path = self.book("typed-groups.xlsx", ["Group", "ID"], rows)
        workbook = openpyxl.load_workbook(path)
        workbook.active["A3"].number_format = "0000"
        workbook.save(path)
        workbook.close()
        result = self.run_job([self.source(path)], mode="split_column", split_column="Group")
        self.assert_manifest(result, 7)
        memberships = {frozenset(row[1][0] for row in self.read_output(path)[1:])
                       for path in result.files}
        self.assertEqual(memberships, {frozenset(("n-one", "n-one-formatted")),
                                      frozenset(("bool",)), frozenset(("text",)),
                                      frozenset(("identifier",)), frozenset(("blank",)),
                                      frozenset(("space",))})

    def test_group_filenames_are_unique_safe_and_manifest_maps_every_file(self):
        labels = ["A/B", "A\\B", "CON", "con", "NUL", "..", "trailing. ",
                  "长" * 200, "SYNTHETIC_GROUP"]
        path = self.book("group-labels.xlsx", ["Group", "ID"], [[v, i] for i, v in enumerate(labels)])
        result = self.run_job([self.source(path)], mode="split_column", split_column="Group")
        manifest = self.assert_manifest(result, len(labels))
        self.assertEqual(len(result.files), len(labels))
        self.assertEqual(len({p.name.casefold() for p in result.files}), len(labels))
        self.assertEqual(Counter(self.value_rows(result)), Counter((v, i) for i, v in enumerate(labels)))
        self.assertEqual(len({entry["group_id"] for entry in manifest["files"]}), len(labels))
        for output in result.files:
            self.assertFalse(any(char in output.name for char in '<>:"/\\|?*'))
            self.assertEqual(output.name, output.name.rstrip(" ."))

    def test_split_requires_single_source_and_valid_options(self):
        one = self.book("one.xlsx", ["ID"], [[1]])
        two = self.book("two.xlsx", ["ID"], [[2]])
        for mode in ("split_rows", "split_column"):
            with self.subTest(mode=mode):
                self.assert_rejected_cleanly([self.source(one), self.source(two)], mode=mode, split_column="ID")
        for count in (0, -1):
            with self.subTest(count=count):
                self.assert_rejected_cleanly([self.source(one)], mode="split_rows", rows_per_file=count)
        self.assert_rejected_cleanly([self.source(one)], mode="split_column", split_column="missing")

    def test_cancel_before_start_and_after_staged_output_leaves_no_published_task(self):
        path = self.book("cancel.xlsx", ["ID"], [[i] for i in range(8)])
        before = digest(path)
        event = threading.Event()
        event.set()
        with self.assertRaises(mod.TableCancelled):
            self.run_job([self.source(path)], mode="split_rows", rows_per_file=2, cancel=event)
        self.assertEqual(list(self.output.iterdir()), [])
        class CancelAfterAFile:
            observed = False
            def is_set(inner):
                inner.observed = inner.observed or any(self.output.rglob("*.xlsx"))
                return inner.observed
        cancel = CancelAfterAFile()
        with self.assertRaises(mod.TableCancelled):
            self.run_job([self.source(path)], mode="split_rows", rows_per_file=2, cancel=cancel)
        self.assertTrue(cancel.observed, "Cancellation must be observed after staging starts")
        self.assertEqual(list(self.output.iterdir()), [])
        self.assertEqual(digest(path), before)

    def test_repeated_runs_preserve_source_previous_task_and_unrelated_file(self):
        path = self.book("repeat.xlsx", ["ID"], [["001"]])
        sentinel = self.output / "unrelated.txt"
        sentinel.write_text("SYNTHETIC existing user file", encoding="utf-8")
        original = digest(path)
        first = self.run_job([self.source(path)], mode="merge")
        previous = {p: digest(p) for p in first.directory.iterdir()}
        second = self.run_job([self.source(path)], mode="merge")
        self.assertNotEqual(first.directory, second.directory)
        self.assertEqual({p: digest(p) for p in first.directory.iterdir()}, previous)
        self.assertEqual(digest(path), original)
        self.assertEqual(sentinel.read_text(encoding="utf-8"), "SYNTHETIC existing user file")
        self.assertEqual(self.value_rows(second), [("001",)])

    def test_published_directory_already_contains_all_valid_outputs_and_manifest(self):
        path = self.book("publish.xlsx", ["ID"], [[i] for i in range(5)])
        real_rename = os.rename
        inspected = []

        def inspect(source, target, *args, **kwargs):
            source, target = Path(source), Path(target)
            if source.is_dir() and target.parent.parent.resolve() == self.output.resolve():
                self.assertFalse(target.exists())
                manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
                self.assertEqual(manifest["rows"], 5)
                self.assertEqual(sum(item["rows"] for item in manifest["files"]), 5)
                files = sorted(source.glob("*.xlsx"))
                self.assertEqual(len(files), 3)
                self.assertEqual({p.name for p in files}, {item["name"] for item in manifest["files"]})
                rows = []
                for output in files:
                    rows.extend(row[0][0] for row in self.read_output(output)[1:])
                self.assertEqual(sorted(rows), list(range(5)))
                inspected.append(target)
            return real_rename(source, target, *args, **kwargs)

        with mock.patch("os.rename", side_effect=inspect):
            result = self.run_job([self.source(path)], mode="split_rows", rows_per_file=2)
        self.assertEqual(inspected, [result.directory])
        self.assert_manifest(result, 5)

    def test_partial_zip_write_enospc_keeps_original_error_and_previous_task(self):
        path = self.book("disk-full.xlsx", ["ID"], [["001"], ["002"]])
        first = self.run_job([self.source(path)], mode="merge")
        before = {p: digest(p) for p in first.directory.iterdir()}
        source_before = digest(path)
        original = ZipFile.open
        failure = OSError(errno.ENOSPC, "SYNTHETIC output disk full")
        observed = []

        class BrokenWriter:
            def __init__(self, stream):
                self.stream = stream
            def __getattr__(self, name):
                return getattr(self.stream, name)
            def __enter__(self):
                return self
            def __exit__(self, *args):
                self.stream.close()
            def write(self, contents):
                self.stream.write(contents[:32])
                observed.append(len(contents[:32]))
                raise failure

        def inject(archive, name, mode="r", *args, **kwargs):
            stream = original(archive, name, mode, *args, **kwargs)
            return BrokenWriter(stream) if mode == "w" else stream

        with mock.patch.object(ZipFile, "open", inject):
            with self.assertRaises(OSError) as raised:
                self.run_job([self.source(path)], mode="split_rows", rows_per_file=1)
        self.assertTrue(observed, "The fault must occur after writing actual ZIP payload")
        self.assertIs(raised.exception, failure)
        self.assertEqual(raised.exception.errno, errno.ENOSPC)
        self.assertEqual(set(self.output.iterdir()), {first.directory.parent})
        self.assertEqual({p: digest(p) for p in first.directory.iterdir()}, before)
        self.assertEqual(digest(path), source_before)

    @unittest.skipUnless(os.name == "nt", "Real read sharing lock requires Windows")
    def test_real_exclusive_input_handle_fails_cleanly_then_retry_succeeds(self):
        import ctypes
        from ctypes import wintypes

        path = self.book("read-locked.xlsx", ["ID"], [["SYNTHETIC"]])
        original = digest(path)
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                      wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE)
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = (wintypes.HANDLE,)
        kernel.CloseHandle.restype = wintypes.BOOL
        owned = kernel.CreateFileW(str(path), 0x80000000, 0, None, 3, 0x80, None)
        self.assertNotEqual(owned, ctypes.c_void_p(-1).value)
        try:
            with self.assertRaises((OSError, ValueError)) as raised:
                self.run_job([self.source(path)], mode="merge")
            self.assertIn(path.name, str(raised.exception))
            self.assertEqual(list(self.output.iterdir()), [])
        finally:
            self.assertTrue(kernel.CloseHandle(owned))
        self.assertEqual(digest(path), original)
        result = self.run_job([self.source(path)], mode="merge")
        self.assertEqual(self.value_rows(result), [("SYNTHETIC",)])

    def test_output_file_limit_rejects_high_cardinality_and_rollover_without_publish(self):
        groups = self.book("too-many-groups.xlsx", ["Group", "ID"],
                           [["group-{}".format(i), i] for i in range(4)])
        with mock.patch.object(mod, "MAX_OUTPUT_FILES", 3):
            self.assert_rejected_cleanly([self.source(groups)], mode="split_column", split_column="Group")
        self.assertEqual(list(self.output.iterdir()), [])
        rows = self.book("too-many-parts.xlsx", ["Group", "ID"], [["same", i] for i in range(4)])
        with mock.patch.object(mod, "MAX_OUTPUT_FILES", 3):
            self.assert_rejected_cleanly([self.source(rows)], mode="split_rows", rows_per_file=1)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_cancellation_during_zip_packaging_cleans_whole_task(self):
        path = self.book("cancel-packaging.xlsx", ["ID"], [[i] for i in range(5)])
        original = ZipFile.open
        stop = threading.Event()
        writes = []

        class CancellingWriter:
            def __init__(self, stream):
                self.stream = stream
            def __getattr__(self, name):
                return getattr(self.stream, name)
            def __enter__(self):
                return self
            def __exit__(self, *args):
                self.stream.close()
            def write(self, contents):
                result = self.stream.write(contents)
                writes.append(len(contents))
                stop.set()
                return result

        def inject(archive, name, mode="r", *args, **kwargs):
            stream = original(archive, name, mode, *args, **kwargs)
            return CancellingWriter(stream) if mode == "w" else stream

        before = digest(path)
        with mock.patch.object(ZipFile, "open", inject):
            with self.assertRaises(mod.TableCancelled):
                self.run_job([self.source(path)], mode="split_rows", rows_per_file=2, cancel=stop)
        self.assertTrue(writes)
        self.assertEqual(list(self.output.iterdir()), [])
        self.assertEqual(digest(path), before)

    def test_source_mtime_change_during_reading_rejects_the_whole_task(self):
        path = self.book("changed-input.xlsx", ["ID"], [[i] for i in range(3)])
        before = path.stat()
        original = digest(path)
        changed = []

        def mutate(detail):
            if detail["phase"] == "writing" and not changed:
                os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns + 2_000_000_000))
                changed.append(True)

        with self.assertRaisesRegex(ValueError, "变化"):
            self.run_job([self.source(path)], mode="merge", progress=mutate)
        self.assertTrue(changed)
        self.assertEqual(list(self.output.iterdir()), [])
        self.assertEqual(digest(path), original)

    def test_publish_permission_failure_cleans_task_and_preserves_same_error(self):
        path = self.book("publish-denied.xlsx", ["ID"], [["SYNTHETIC"]])
        failure = PermissionError(errno.EACCES, "SYNTHETIC publish denied")
        with mock.patch("os.rename", side_effect=failure):
            with self.assertRaises(PermissionError) as raised:
                self.run_job([self.source(path)], mode="merge")
        self.assertIs(raised.exception, failure)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_cleanup_failure_preserves_errno_and_identifies_only_own_stage(self):
        path = self.book("cleanup-locked.xlsx", ["ID"], [["SYNTHETIC"]])
        failure = OSError(errno.ENOSPC, "SYNTHETIC disk full with cleanup failure")
        cleanup = PermissionError(errno.EACCES, "SYNTHETIC cleanup lock")
        original_open = ZipFile.open
        original_rmtree = mod.shutil.rmtree
        stages = []

        def writing_failure(archive, name, mode="r", *args, **kwargs):
            if mode == "w":
                raise failure
            return original_open(archive, name, mode, *args, **kwargs)

        def locked_stage(target, *args, **kwargs):
            target = Path(target)
            if target.parent.resolve() == self.output.resolve():
                stages.append(target)
                raise cleanup
            return original_rmtree(target, *args, **kwargs)

        with mock.patch.object(ZipFile, "open", writing_failure), mock.patch.object(
                mod.shutil, "rmtree", side_effect=locked_stage):
            with self.assertRaises(OSError) as raised:
                self.run_job([self.source(path)], mode="merge")
        self.assertIs(raised.exception, failure)
        self.assertEqual(raised.exception.errno, errno.ENOSPC)
        self.assertEqual(len(stages), 1)
        self.assertEqual(Path(failure.partial_path), stages[0])
        self.assertIn("SYNTHETIC cleanup lock", str(failure.cleanup_error))
        self.assertTrue(stages[0].is_dir())
        self.assertFalse(any(self.output.glob("*/结果")))
        details = str(failure) + "\n".join(getattr(failure, "__notes__", []))
        self.assertIn(str(stages[0]), details)

    def test_cancel_interrupts_sqlite_index_and_cleans_grouping_task(self):
        path = self.book("index-cancel.xlsx", ["Group", "ID"], [[i % 2, i] for i in range(5000)])
        cancel = threading.Event()
        connect = sqlite3.connect
        interrupted = []

        class Connection:
            def __init__(self, connection):
                self.connection = connection
            def __getattr__(self, name):
                return getattr(self.connection, name)
            def execute(self, statement, *args, **kwargs):
                if statement.lstrip().upper().startswith("CREATE INDEX"):
                    cancel.set()
                    try:
                        return self.connection.execute(statement, *args, **kwargs)
                    except sqlite3.OperationalError as error:
                        interrupted.append(str(error))
                        raise
                return self.connection.execute(statement, *args, **kwargs)

        with mock.patch.object(sqlite3, "connect", side_effect=lambda *a, **k: Connection(connect(*a, **k))):
            with self.assertRaises(mod.TableCancelled):
                self.run_job([self.source(path)], mode="split_column", split_column="Group", cancel=cancel)
        self.assertTrue(interrupted, "The SQL index execution, not only the following loop, must be interrupted")
        self.assertEqual(list(self.output.iterdir()), [])

    def test_concurrent_tasks_reserve_distinct_parents_without_overwriting(self):
        path = self.book("parallel.xlsx", ["ID"], [["SYNTHETIC"]])
        sentinel = self.output / "合表_SYNTHETIC_FIXED"
        sentinel.write_bytes(b"SYNTHETIC existing collision")
        original = digest(path)
        barrier = threading.Barrier(2)

        def at_publish(detail):
            if detail["phase"] == "publishing":
                barrier.wait(timeout=5)

        with mock.patch.object(mod.time, "strftime", return_value="SYNTHETIC_FIXED"):
            with ThreadPoolExecutor(max_workers=2) as executor:
                futures = [executor.submit(self.run_job, [self.source(path)], mode="merge", progress=at_publish)
                           for _ in range(2)]
                results = [future.result(timeout=15) for future in futures]
        self.assertEqual(len({result.directory for result in results}), 2)
        for result in results:
            self.assert_manifest(result, 1)
            self.assertEqual(self.value_rows(result), [("SYNTHETIC",)])
        self.assertEqual(sentinel.read_bytes(), b"SYNTHETIC existing collision")
        self.assertEqual(digest(path), original)


if __name__ == "__main__":
    unittest.main()
