import hashlib, importlib.util, tempfile, unittest
from datetime import datetime
from contextlib import closing
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape, quoteattr
from xml.etree import ElementTree as ET
import openpyxl

SCRIPT = Path(__file__).resolve().parents[1] / "excel_unmerge_fill.py"
spec = importlib.util.spec_from_file_location('unmerge', SCRIPT)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'

def c(ref, value=None, kind='str', style=None):
    attrs=' r='+quoteattr(ref)+((' s='+quoteattr(str(style))) if style is not None else '')
    if kind=='str': return '<c'+attrs+' t="inlineStr"><is><t xml:space="preserve">'+escape(str(value))+'</t></is></c>'
    if kind=='blank': return '<c'+attrs+'/>'
    if kind=='formula': return '<c'+attrs+'><f>'+escape(str(value))+'</f><v>9</v></c>'
    if kind=='bool': return '<c'+attrs+' t="b"><v>'+str(int(value))+'</v></c>'
    if kind=='shared': return '<c'+attrs+' t="s"><v>'+str(value)+'</v></c>'
    return '<c'+attrs+'><v>'+str(value)+'</v></c>'

def sheet(rows, merges=(), extra=''):
    xml='<worksheet xmlns="'+NS+'"><dimension ref="A1:K20"/><sheetData>'
    for num, cells in sorted(rows.items()): xml+='<row r="'+str(num)+'">'+''.join(cells)+'</row>'
    xml+='</sheetData>'
    if merges: xml+='<mergeCells count="'+str(len(merges))+'">'+''.join('<mergeCell ref='+quoteattr(m)+'/>' for m in merges)+'</mergeCells>'
    return (xml+extra+'</worksheet>').encode()

STYLES=('<styleSheet xmlns="'+NS+'"><fonts count="1"><font><name val="Calibri"/><sz val="11"/></font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>').encode()

def makebook(path, sheets, extra_parts=None):
    wbtype='application/vnd.ms-excel.sheet.macroEnabled.main+xml' if path.suffix=='.xlsm' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'
    types='<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="bin" ContentType="application/octet-stream"/><Override PartName="/xl/workbook.xml" ContentType="'+wbtype+'"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
    workbook='<workbook xmlns="'+NS+'" xmlns:r="'+REL+'"><sheets>'
    rels='<Relationships xmlns="'+PKG+'"><Relationship Id="styles" Type="'+REL+'/styles" Target="styles.xml"/><Relationship Id="strings" Type="'+REL+'/sharedStrings" Target="sharedStrings.xml"/>'
    parts={}
    for i,(name,raw) in enumerate(sheets,1):
        filename='xl/worksheets/sheet'+str(i)+'.xml'
        types+='<Override PartName="/'+filename+'" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        workbook+='<sheet name='+quoteattr(name)+' sheetId="'+str(i)+'" r:id="rId'+str(i)+'"'+(' state="hidden"' if name.startswith('隐藏') else '')+'/>'
        rels+='<Relationship Id="rId'+str(i)+'" Type="'+REL+'/worksheet" Target="/xl/worksheets/sheet'+str(i)+'.xml"/>'
        parts[filename]=raw
    parts.update({'[Content_Types].xml':types+'</Types>','_rels/.rels':'<Relationships xmlns="'+PKG+'"><Relationship Id="wb" Type="'+REL+'/officeDocument" Target="xl/workbook.xml"/></Relationships>','xl/workbook.xml':workbook+'</sheets></workbook>','xl/_rels/workbook.xml.rels':rels+'</Relationships>','xl/styles.xml':STYLES,'xl/sharedStrings.xml':'<sst xmlns="'+NS+'" count="1" uniqueCount="1"><si><t>共享内容</t></si></sst>','opaque/sentinel.bin':b'untouched\x00\xff'})
    parts.update(extra_parts or {})
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        z.comment=b'keep archive comment'
        for k,v in parts.items(): z.writestr(k,v)
    return path

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def xml(path, part='xl/worksheets/sheet1.xml'):
    with ZipFile(path) as z: return ET.fromstring(z.read(part))
def merge_refs(path,part='xl/worksheets/sheet1.xml'):
    return [n.attrib['ref'] for n in xml(path,part).findall('{'+NS+'}mergeCells/{'+NS+'}mergeCell')]
def load(path): return closing(openpyxl.load_workbook(path,read_only=True,data_only=False))

class UnmergeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='excel-unmerge-review-')
        self.root=Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def book(self, rows, merges=(), name='sample.xlsx'):
        return makebook(self.root/name,[('数据',sheet(rows,merges))])
    def assert_unchanged_other_parts(self, src, out, changed):
        with ZipFile(src) as a, ZipFile(out) as b:
            self.assertEqual(a.namelist(),b.namelist())
            self.assertEqual(a.comment,b.comment)
            for n in a.namelist():
                if n not in changed: self.assertEqual(a.read(n),b.read(n),n)
    def test_vertical_and_original_not_modified(self):
        p=self.book({2:[c('A2','华东仓'),c('B2',5,'num')],3:[c('B3',6,'num')]},['A2:A4'])
        digest=sha(p); out,stats=mod.process_file(p)
        with load(out) as wb:
            self.assertEqual([wb.active.cell(r,1).value for r in range(2,5)],['华东仓']*3)
            self.assertEqual(wb.active['B3'].value,6)
        self.assertEqual(sha(p),digest); self.assertEqual(stats,[('数据',1,2)])
        self.assertEqual(merge_refs(out),[])
        self.assert_unchanged_other_parts(p,out,{'xl/worksheets/sheet1.xml'})
    def test_rectangle_preserves_horizontal_and_header(self):
        p=self.book({1:[c('A1','表头')],2:[c('B2','组')]},['A1:D1','B2:D4'])
        out,_=mod.process_file(p)
        self.assertEqual(merge_refs(out),['A1:D1','B2:D2','B3:D3','B4:D4'])
        with load(out) as wb:
            for r in range(2,5): self.assertEqual(wb.active.cell(r,2).value,'组')
            self.assertIsNone(wb.active['C3'].value)
    def test_all_merges(self):
        p=self.book({1:[c('A1','表头')],2:[c('B2','组')]},['A1:D1','B2:D4'])
        out,stats=mod.process_file(p,all_merges=True)
        self.assertEqual(merge_refs(out),[])
        with load(out) as wb:
            self.assertEqual([wb.active.cell(1,col).value for col in range(1,5)],['表头']*4)
            self.assertEqual([wb.active.cell(r,col).value for r in range(2,5) for col in range(2,5)],['组']*9)
        self.assertEqual(stats,[('数据',2,11)])
    def test_types_and_ordinary_blanks(self):
        p=self.book({2:[c('A2',0,'num'),c('B2',False,'bool'),c('C2','0012'),c('D2',45292,'num',1),c('E2',0,'shared'),c('F2','  空格  '),c('H2',8,'num')],3:[c('G3',None,'blank')]},['A2:A3','B2:B3','C2:C3','D2:D3','E2:E3','F2:F3'])
        out,_=mod.process_file(p)
        with load(out) as wb:
            ws=wb.active
            self.assertEqual(ws['A3'].value,0); self.assertIs(ws['B3'].value,False)
            self.assertEqual(ws['C3'].value,'0012'); self.assertEqual(ws['D3'].value,datetime(2024,1,1))
            self.assertEqual(ws['E3'].value,'共享内容'); self.assertEqual(ws['F3'].value,'  空格  ')
            self.assertIsNone(ws['G3'].value); self.assertIsNone(ws['H3'].value)
    def test_empty_anchor(self):
        p=self.book({2:[c('A2',None,'blank')]},['A2:A5'])
        out,stats=mod.process_file(p)
        self.assertEqual(stats,[('数据',1,0)])
        with load(out) as wb: self.assertTrue(all(wb.active.cell(r,1).value is None for r in range(2,6)))
        self.assertEqual(merge_refs(out),[])
    def test_no_matches_no_output(self):
        for i,merges in enumerate([[],['A1:C1']]):
            p=self.book({1:[c('A1','表头')]},merges,'no'+str(i)+'.xlsx')
            before=set(self.root.iterdir()); out,stats=mod.process_file(p)
            self.assertIsNone(out); self.assertEqual(set(self.root.iterdir()),before)
            self.assertEqual(stats,[('数据',0,0)])
    def test_non_overwrite(self):
        p=self.book({1:[c('A1','x')]},['A1:A2'])
        out1,_=mod.process_file(p); digest=sha(out1)
        out2,_=mod.process_file(p)
        self.assertEqual(out1.name,'sample_拆分填充.xlsx')
        self.assertEqual(out2.name,'sample_拆分填充_2.xlsx')
        self.assertEqual(sha(out1),digest)
        with load(out2) as wb: self.assertEqual(wb.active['A2'].value,'x')
    def test_multi_sheet_selection_and_hidden(self):
        p=makebook(self.root/'multi.xlsx',[('明细',sheet({1:[c('A1','one')]},['A1:A2'])),('隐藏辅助',sheet({1:[c('A1','two')]},['A1:A3']))])
        out,stats=mod.process_file(p,sheets=['隐藏辅助'])
        self.assertEqual(stats,[('隐藏辅助',1,2)])
        self.assert_unchanged_other_parts(p,out,{'xl/worksheets/sheet2.xml'})
        with load(out) as wb:
            self.assertEqual(wb['隐藏辅助']['A3'].value,'two'); self.assertEqual(wb['隐藏辅助'].sheet_state,'hidden')
        both,stats=mod.process_file(p)
        self.assertEqual(len(stats),2)
        with load(both) as wb: self.assertEqual(wb['明细']['A2'].value,'one')
    def test_formula_rejects_whole_workbook_without_output(self):
        p=makebook(self.root/'formula.xlsx',[('有效',sheet({1:[c('A1','x')]},['A1:A2'])),('公式',sheet({1:[c('A1','1+8','formula')]},['A1:A2']))])
        before=set(self.root.iterdir()); digest=sha(p)
        with self.assertRaisesRegex(ValueError,'公式'): mod.process_file(p)
        self.assertEqual(set(self.root.iterdir()),before); self.assertEqual(sha(p),digest)
    def test_unrelated_formula_preserved(self):
        p=self.book({1:[c('A1','x'),c('B1','1+8','formula')]},['A1:A2'])
        out,_=mod.process_file(p)
        before=xml(p).find('.//{'+NS+'}c[@r="B1"]'); after=xml(out).find('.//{'+NS+'}c[@r="B1"]')
        self.assertEqual(ET.tostring(before),ET.tostring(after))
        with load(out) as wb: self.assertEqual(wb.active['B1'].value,'=1+8')
    def test_conflicting_data_rejected(self):
        p=self.book({1:[c('A1','x')],2:[c('A2','hidden')]},['A1:A2'])
        before=set(self.root.iterdir())
        with self.assertRaisesRegex(ValueError,'A2'): mod.process_file(p)
        self.assertEqual(set(self.root.iterdir()),before)
    def test_xlsm_preserves_other_bytes(self):
        p=makebook(self.root/'macro.xlsm',[('宏表',sheet({1:[c('A1','x')]},['A1:A2']))],{'xl/vbaProject.bin':b'macro-content-preservation-probe'})
        out,_=mod.process_file(p)
        self.assertEqual(out.suffix,'.xlsm')
        self.assert_unchanged_other_parts(p,out,{'xl/worksheets/sheet1.xml'})
        with load(out) as wb: self.assertEqual(wb.active['A2'].value,'x')
    def test_unknown_sheet_no_output(self):
        p=self.book({1:[c('A1','x')]},['A1:A2']); before=set(self.root.iterdir())
        with self.assertRaisesRegex(ValueError,'找不到工作表'): mod.process_file(p,sheets=['不存在'])
        self.assertEqual(set(self.root.iterdir()),before)
    def test_empty_value_target_is_fillable(self):
        p=self.book({1:[c('A1',0,'num')],2:['<c r="A2"><v/></c>']},['A1:A2'])
        out,stats=mod.process_file(p)
        with load(out) as wb: self.assertEqual(wb.active['A2'].value,0)
        self.assertEqual(stats,[('数据',1,1)])
    def test_row_extension_order(self):
        raw=sheet({1:[c('A1','x'),'<extLst><ext uri="test"/></extLst>']},['A1:A2'])
        out,_,_=mod.transform_sheet(raw)
        row=ET.fromstring(out).find('{'+NS+'}sheetData/{'+NS+'}row')
        names=[n.tag.split('}')[-1] for n in row]
        self.assertEqual(names,['c','extLst'])

if __name__=='__main__': unittest.main(verbosity=2)
