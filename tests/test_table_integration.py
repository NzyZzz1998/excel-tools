"""Real Tk parent/Toplevel integration; only synthetic, owned windows/jobs.

All Tk teardown and cyclic collection happen on the interpreter's creator
thread before subsequent tests can start background workers.
"""
import gc
from pathlib import Path
import tempfile
import threading
import time
import tkinter as tk
import unittest
import weakref
from unittest import mock

from openpyxl import Workbook

import excel_table_gui as table_gui
from excel_unmerge_gui import Application


class TableRootIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = Application(self.root)
        self.child = None
        self.workers = []
        self.releases = []
        self.callback_errors = []
        errors = self.callback_errors
        self.root.report_callback_exception = lambda kind, value, tb: errors.append((kind.__name__, str(value)))
        self.temporary = tempfile.TemporaryDirectory(prefix="excel-table-integration-")
        self.folder = Path(self.temporary.name)
        self.root.update()

    def alive(self):
        try:
            return bool(self.root.winfo_exists())
        except tk.TclError:
            return False

    def wait_for(self, condition):
        deadline = time.monotonic() + 10
        while not condition() and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertTrue(condition(), "Owned root/child did not reach the expected state")

    def tearDown(self):
        self.assertIs(threading.current_thread(), threading.main_thread())
        for release in self.releases:
            release.set()
        if self.child is not None and self.child.running:
            self.child.request_cancel()
            self.wait_for(lambda: not self.child.running)
        if self.app.running:
            self.app.request_stop()
            self.wait_for(lambda: not self.app.running)
        for worker in self.workers:
            worker.join(timeout=5)
            self.assertFalse(worker.is_alive(), "Owned worker remained alive during Tk disposal")
        references = [weakref.ref(self.app), weakref.ref(self.root)]
        if self.child is not None:
            references.extend((weakref.ref(self.child), weakref.ref(self.child.window)))
            if not self.child.closed:
                self.child.close()
        if self.alive():
            self.root.destroy()
        self.child = None
        self.app = None
        self.root = None
        self.workers.clear()
        gc.collect()
        self.temporary.cleanup()
        self.assertEqual(self.callback_errors, [])
        for reference in references:
            self.assertIsNone(reference(), "Destroyed Tk object was retained past main-thread collection")

    @staticmethod
    def disabled(widget):
        return str(widget.cget("state")) == "disabled"

    def open_child(self):
        self.app.open_table_tools()
        self.child = self.app.table_tools
        self.assertIsNotNone(self.child)
        # Hide before processing geometry events; never bring a test window
        # over a user's application or manipulate any pre-existing Toplevel.
        self.child.window.attributes("-alpha", 0)
        self.child.window.withdraw()
        self.root.update()
        return self.child

    def prepare_child(self):
        child = self.open_child()
        path = self.folder / "synthetic.xlsx"
        workbook = Workbook()
        workbook.active.title = "数据"
        workbook.active.append(["ID", "Value"])
        workbook.active.append(["001", 1])
        workbook.active.append(["002", 2])
        workbook.save(path)
        workbook.close()
        child.sources = [table_gui.tables.TableSource(path, "数据", 1)]
        child._sources_changed()
        child.output_dir.set(str(self.folder))
        child.preview()
        if child.worker is not None:
            self.workers.append(child.worker)
        self.wait_for(lambda: not child.running)
        self.assertFalse(self.disabled(child.run_button))
        return child

    def seed_previous_task(self):
        self.app.files = (str(self.folder / "next-choice.xlsx"),)
        self.app.file_states = {
            str(self.folder / "old-success.xlsx"): "success",
            str(self.folder / "old-failed.xlsx"): "failed",
            str(self.folder / "old-pending.xlsx"): "pending",
        }
        self.app.batch_all_merges = True
        self.app.all_merges.set(True)
        self.app.show_next_selection()
        self.app.update_counts()

    def test_open_child_disables_editing_and_all_legacy_entry_points(self):
        self.seed_previous_task()
        selection = self.app.files
        states = dict(self.app.file_states)
        child = self.open_child()
        for widget in (self.app.choose_button, self.app.start_button, self.app.retry_button,
                       self.app.resume_button, self.app.option):
            self.assertTrue(self.disabled(widget))
        with mock.patch.object(self.app, "launch_batch") as launch:
            self.app.start()
            self.app.retry_failed()
            self.app.resume_pending()
            launch.assert_not_called()
        self.assertFalse(self.app.running)
        self.assertIs(self.app.table_tools, child)
        self.assertEqual(self.app.files, selection)
        self.assertEqual(self.app.file_states, states)
        self.assertTrue(self.app.all_merges.get())

    def test_child_close_restores_buttons_for_actual_previous_and_next_states(self):
        cases = [([], {}), (["next.xlsx"], {}),
                 ([], {"failed.xlsx": "failed"}),
                 (["next.xlsx"], {"success.xlsx": "success", "pending.xlsx": "pending"}),
                 (["next.xlsx"], {"failed.xlsx": "failed", "pending.xlsx": "pending"})]
        for files, states in cases:
            with self.subTest(files=files, states=states):
                self.app.files = tuple(files)
                self.app.file_states = dict(states)
                child = self.open_child()
                child.close()
                self.root.update()
                self.assertIsNone(self.app.table_tools)
                self.assertTrue(child.closed)
                self.assertFalse(self.disabled(self.app.choose_button))
                self.assertFalse(self.disabled(self.app.option))
                self.assertEqual(self.disabled(self.app.start_button), not bool(files))
                self.assertEqual(self.disabled(self.app.retry_button), "failed" not in states.values())
                self.assertEqual(self.disabled(self.app.resume_button), "pending" not in states.values())
                self.assertEqual(self.app.file_states, states)
                self.assertEqual(self.app.files, tuple(files))

    def test_legacy_worker_blocks_child_until_its_done_event(self):
        entered, release = threading.Event(), threading.Event()
        self.releases.append(release)
        self.app.files = (str(self.folder / "legacy-synthetic.xlsx"),)
        workers = self.workers

        def slow_legacy(file, all_merges=False, progress=None):
            workers.append(threading.current_thread())
            entered.set()
            if not release.wait(timeout=5):
                raise AssertionError("Owned legacy test worker was not released")
            return None, []

        with mock.patch("excel_unmerge_gui.process_file", side_effect=slow_legacy):
            self.app.start()
            try:
                self.wait_for(entered.is_set)
                self.assertTrue(self.disabled(self.app.tools_button))
                self.app.open_table_tools()
                self.assertIsNone(self.app.table_tools)
            finally:
                release.set()
                self.wait_for(lambda: not self.app.running)
        self.assertFalse(self.disabled(self.app.tools_button))
        self.open_child()

    def test_root_close_only_requests_child_cancel_and_keeps_both_windows_alive(self):
        self.seed_previous_task()
        states = dict(self.app.file_states)
        child = self.prepare_child()
        entered, release = threading.Event(), threading.Event()
        self.releases.append(release)

        def slow_table(*, cancel, **kwargs):
            entered.set()
            if not release.wait(timeout=5):
                raise AssertionError("Owned table test worker was not released")
            if cancel.is_set():
                raise table_gui.tables.TableCancelled("SYNTHETIC cancelled after cleanup")
            raise AssertionError("The root must request cooperative child cancellation")

        with mock.patch.object(table_gui.tables, "run_tables", side_effect=slow_table):
            child.start()
            self.workers.append(child.worker)
            try:
                self.wait_for(entered.is_set)
                self.app.close()
                self.app.close()
                self.assertTrue(child.cancel_event.is_set())
                self.assertTrue(child.running)
                self.assertFalse(child.closed)
                self.assertTrue(self.alive())
                self.assertTrue(child.window.winfo_exists())
                with mock.patch.object(self.app, "launch_batch") as launch:
                    self.app.start()
                    self.app.retry_failed()
                    self.app.resume_pending()
                    launch.assert_not_called()
            finally:
                release.set()
                self.wait_for(lambda: not child.running)
        self.assertTrue(self.alive())
        self.assertEqual(self.app.file_states, states)
        child.close()
        self.assertIsNone(self.app.table_tools)
        self.assertFalse(self.disabled(self.app.retry_button))
        self.assertFalse(self.disabled(self.app.resume_button))

    def test_repeated_entry_raises_existing_toplevel_without_constructing_another(self):
        child = self.open_child()
        before = [widget for widget in self.root.winfo_children() if isinstance(widget, tk.Toplevel)]
        with mock.patch.object(table_gui, "TableToolsWindow", side_effect=AssertionError("Duplicate child")):
            # Observe focus/restore intent without activating a hidden test
            # window over a user's foreground application.
            with mock.patch.object(child.window, "lift") as raised, mock.patch.object(
                    child.window, "deiconify") as restored, mock.patch.object(
                    child.window, "focus_set") as focused:
                self.app.open_table_tools()
                raised.assert_called_once()
                restored.assert_called_once()
                focused.assert_called_once()
        self.assertIs(self.app.table_tools, child)
        after = [widget for widget in self.root.winfo_children() if isinstance(widget, tk.Toplevel)]
        self.assertEqual(after, before)

    def test_completed_child_can_close_root_and_leaves_valid_output(self):
        child = self.prepare_child()
        child.start()
        self.workers.append(child.worker)
        self.wait_for(lambda: not child.running)
        self.assertIsNotNone(child.result)
        self.assertEqual(child.result.rows, 2)
        self.assertTrue(child.result.manifest.is_file())
        self.assertTrue(all(path.is_file() for path in child.result.files))
        self.app.close()
        self.assertTrue(child.closed)
        self.assertIsNone(self.app.table_tools)
        self.assertFalse(self.alive())

    def test_closed_child_is_collected_on_main_thread_before_reopening(self):
        child = self.open_child()
        child_ref, window_ref = weakref.ref(child), weakref.ref(child.window)
        child.close()
        self.assertIsNone(child._after_id)
        self.assertIsNone(child.on_close)
        self.assertIsNone(self.app.table_tools)
        self.child = None
        del child
        self.assertIs(threading.current_thread(), threading.main_thread())
        gc.collect()
        self.assertIsNone(child_ref(), "Parent retained the closed table window")
        self.assertIsNone(window_ref(), "Closed Toplevel retained its Tcl/Tk variables")
        self.open_child()


if __name__ == "__main__":
    unittest.main()
