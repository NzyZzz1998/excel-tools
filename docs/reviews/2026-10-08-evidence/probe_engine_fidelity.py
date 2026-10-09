import hashlib
import json
import struct
import subprocess
import sys
import zlib
from contextlib import closing
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import openpyxl
import xlsxwriter

REPO = Path(r'E:\codex\excel-tools')
sys.path.insert(0, str(REPO))
import excel_unmerge_fill as engine

ROOT = Path(__file__).resolve().parent
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def chunk(kind, data):
    return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
image = (b'\x89PNG\r\n\x1a\n'
         + chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
         + chunk(b'IDAT', zlib.compress(b'\x00\xff\x00\x00'))
         + chunk(b'IEND', b''))
image_path = ROOT / 'one-red-pixel.png'
image_path.write_bytes(image)
source = ROOT / 'merged-in-cell-image.xlsx'
with xlsxwriter.Workbook(source) as book:
    sheet = book.add_worksheet()
    sheet.merge_range('A1:A2', '')
    embed_return = sheet.embed_image('A1', image_path)
output, stats = engine.process_file(source)
with ZipFile(source) as original, ZipFile(output) as transformed:
    before = ET.fromstring(original.read('xl/worksheets/sheet1.xml'))
    after = ET.fromstring(transformed.read('xl/worksheets/sheet1.xml'))
    rich_parts = [p for p in original.namelist() if 'richData' in p or 'metadata' in p or 'media/' in p]
    part_bytes_preserved = all(original.read(p) == transformed.read(p) for p in rich_parts)
    ROOT.joinpath('input-sheet.xml').write_bytes(original.read('xl/worksheets/sheet1.xml'))
    ROOT.joinpath('output-sheet.xml').write_bytes(transformed.read('xl/worksheets/sheet1.xml'))
    def cells(root):
        return [{'attributes': c.attrib, 'value': c.find('{'+NS+'}v').text}
                for c in root.findall('.//{'+NS+'}c')]
    result = {
        'repo': str(REPO),
        'git_head': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=REPO, text=True).strip(),
        'source_file_sha256': hashlib.sha256((REPO/'excel_unmerge_fill.py').read_bytes()).hexdigest(),
        'python_version': sys.version,
        'xlsxwriter_version': xlsxwriter.__version__,
        'openpyxl_version': openpyxl.__version__,
        'embedded_image': {
            'input': str(source), 'output': str(output),
            'embed_return': embed_return, 'stats': stats,
            'before_cells': cells(before), 'after_cells': cells(after),
            'rich_parts': rich_parts, 'rich_parts_bytes_unchanged': part_bytes_preserved,
            'excel365_ui_opened': False,
        },
        'omitted_coordinate_probes': [],
    }
normal = ROOT/'normal.xlsx'
book = openpyxl.Workbook()
sheet = book.active
sheet['A1'] = 'VALUE'
sheet.merge_cells('A1:A2')
book.save(normal)
book.close()
for probe in ('row_without_r', 'cell_without_r'):
    special = ROOT/(probe+'.xlsx')
    with ZipFile(normal) as original, ZipFile(special, 'w', ZIP_DEFLATED) as transformed:
        for info in original.infolist():
            payload = original.read(info)
            if info.filename == 'xl/worksheets/sheet1.xml':
                if probe == 'row_without_r':
                    payload = payload.replace(b'<row r="1"', b'<row', 1)
                else:
                    payload = payload.replace(b'<c r="A1"', b'<c', 1)
            transformed.writestr(info, payload)
    with closing(openpyxl.load_workbook(special, read_only=True)) as book:
        reader_value = book.active['A1'].value
    try:
        engine.process_file(special)
        engine_result = 'accepted'
    except Exception as error:
        engine_result = type(error).__name__ + ': ' + str(error)
    result['omitted_coordinate_probes'].append({
        'probe': probe, 'input': str(special),
        'openpyxl_a1': reader_value, 'engine_result': engine_result,
    })
result_path = ROOT/'results.json'
result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'evidence_dir': str(ROOT), 'result': result}, ensure_ascii=True, indent=2))
