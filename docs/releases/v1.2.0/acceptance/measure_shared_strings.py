"""Pure synthetic large-SST acceptance probe; not representative business data.

Run with Python 3.12 plus psutil and openpyxl test dependencies. Only the worker's RSS is measured;
fixture generation and the independent openpyxl oracle run outside that window.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from zipfile import ZIP_DEFLATED, ZipFile


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "excel_table_io.py").is_file())
HERE = ROOT / "testfile" / "v1.2.0-sst-stress"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
ROWS = 100_000
STRING_BYTES = 544


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def payload(index):
    prefix = f"SYNTHETIC-{index:06d}|"
    seed = hashlib.sha256(f"sst-stress-{index}".encode("ascii")).hexdigest()
    return (prefix + seed * 9)[:STRING_BYTES]


def digest_row(digest, index, text):
    digest.update(json.dumps([index, text], ensure_ascii=True, separators=(",", ":")).encode("ascii"))
    digest.update(b"\n")


def create_fixture(path):
    """Independent OOXML fixture; never uses the application's writer."""
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '</Types>'),
        "_rels/.rels": f'<Relationships xmlns="{PKG}"><Relationship Id="wb" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{NS}" xmlns:r="{REL}"><sheets><sheet name="Synthetic SST" sheetId="1" r:id="sheet"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{PKG}"><Relationship Id="sheet" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="sst" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
    }
    expected = hashlib.sha256()
    with path.open("xb") as target, ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, contents in parts.items():
            archive.writestr(name, contents)
        with archive.open("xl/sharedStrings.xml", "w", force_zip64=True) as strings:
            strings.write(f'<sst xmlns="{NS}" count="{ROWS}" uniqueCount="{ROWS}">'.encode())
            for index in range(1, ROWS + 1):
                text = payload(index)
                strings.write(f"<si><t>{text}</t></si>".encode("ascii"))
                digest_row(expected, index, text)
            strings.write(b"</sst>")
        with archive.open("xl/worksheets/sheet1.xml", "w", force_zip64=True) as sheet:
            sheet.write((f'<worksheet xmlns="{NS}"><dimension ref="A1:B{ROWS + 1}"/><sheetData>'
                         '<row r="1"><c r="A1" t="inlineStr"><is><t>Sequence</t></is></c>'
                         '<c r="B1" t="inlineStr"><is><t>Payload</t></is></c></row>').encode())
            for index in range(1, ROWS + 1):
                row = index + 1
                sheet.write(f'<row r="{row}"><c r="A{row}" t="n"><v>{index}</v></c><c r="B{row}" t="s"><v>{index - 1}</v></c></row>'.encode())
            sheet.write(b"</sheetData></worksheet>")
    with ZipFile(path) as archive:
        info = archive.getinfo("xl/sharedStrings.xml")
        sheet_info = archive.getinfo("xl/worksheets/sheet1.xml")
    assert info.file_size > 50 * 1024 * 1024
    return dict(rows=ROWS, columns=2, unique_strings=ROWS, string_bytes=STRING_BYTES,
                sst_uncompressed_bytes=info.file_size, sst_compressed_bytes=info.compress_size,
                sheet_uncompressed_bytes=sheet_info.file_size, source_bytes=path.stat().st_size,
                expected_row_sha256=expected.hexdigest())


def worker(source, output, result_path):
    sys.path.insert(0, str(ROOT))
    from excel_table_tools import TableSource, run_tables
    notifications = 0

    def progress(detail):
        nonlocal notifications
        notifications += 1

    started = time.perf_counter()
    result = run_tables([TableSource(source, "Synthetic SST")], output,
                        mode="split_rows", rows_per_file=40_000,
                        provenance=False, progress=progress)
    seconds = time.perf_counter() - started
    result_path.write_text(json.dumps(dict(rows=result.rows, files=[str(path) for path in result.files],
                                         directory=str(result.directory), wall_seconds=seconds,
                                         progress_events=notifications), indent=2), encoding="utf-8")


def verify_openpyxl(source, files, expected_sha):
    import openpyxl

    def scan(paths, expected_counts):
        digest = hashlib.sha256()
        total = 0
        counts = []
        for path, wanted_count in zip(paths, expected_counts, strict=True):
            book = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                assert len(book.worksheets) == 1
                sheet = book.worksheets[0]
                assert sheet.max_row == wanted_count + 1 and sheet.max_column == 2
                rows = sheet.iter_rows()
                headers = next(rows)
                assert [cell.value for cell in headers] == ["Sequence", "Payload"]
                count = 0
                for row in rows:
                    total += 1
                    count += 1
                    assert len(row) == 2
                    assert row[0].data_type == "n" and row[0].value == total
                    assert row[1].data_type == "s" and row[1].value == payload(total)
                    digest_row(digest, row[0].value, row[1].value)
                assert count == wanted_count
                counts.append(count)
            finally:
                book.close()
        assert total == ROWS and digest.hexdigest() == expected_sha
        return dict(rows=total, counts=counts, ordered_rows_sha256=digest.hexdigest())

    return dict(openpyxl_version=openpyxl.__version__,
                source=scan([source], [ROWS]),
                output=scan(files, [40_000, 40_000, 20_000]))


def main():
    import psutil
    run = HERE / time.strftime("run-%Y%m%d-%H%M%S")
    run.mkdir(parents=True)
    output = run / "outputs"
    output.mkdir()
    source = run / "Synthetic large SST.xlsx"
    fixture = create_fixture(source)
    before = sha(source)
    worker_result = run / "worker-result.json"
    peak_rss = peak_private = samples = 0
    stop_reason = None
    started = time.perf_counter()
    with (run / "worker-stdout.txt").open("w", encoding="utf-8") as stdout, (run / "worker-stderr.txt").open("w", encoding="utf-8") as stderr:
        child = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), "--worker", str(source), str(output), str(worker_result)],
                                 cwd=ROOT, stdout=stdout, stderr=stderr,
                                 creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        process = psutil.Process(child.pid)
        while child.poll() is None:
            try:
                memory = process.memory_info()
                peak_rss = max(peak_rss, memory.rss)
                peak_private = max(peak_private, getattr(memory, "private", 0))
                samples += 1
            except psutil.NoSuchProcess:
                break
            if peak_rss > 1024 ** 3:
                stop_reason = "probe_rss_cap_1GiB"
            if time.perf_counter() - started > 300:
                stop_reason = "probe_timeout_300s"
            if stop_reason:
                child.kill()  # Only this probe's retained Popen process handle.
                break
            time.sleep(.1)
        returncode = child.wait(timeout=15)
    result = dict(fixture=fixture, source_sha256_before=before,
                  source_sha256_after=sha(source), engine_exit_code=returncode,
                  worker_envelope_seconds=time.perf_counter() - started,
                  peak_rss_bytes=peak_rss, peak_rss_MiB=peak_rss / 1024 ** 2,
                  peak_private_bytes=peak_private, sample_interval_seconds=.1,
                  samples=samples, probe_stop_reason=stop_reason,
                  python=sys.version, psutil_version=psutil.__version__,
                  app_sha256={name: sha(ROOT / name) for name in ("excel_table_io.py", "excel_table_tools.py")})
    (run / "measurement.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    assert stop_reason is None and returncode == 0, result
    result["engine"] = json.loads(worker_result.read_text(encoding="utf-8"))
    files = [Path(path) for path in result["engine"]["files"]]
    assert result["engine"]["rows"] == ROWS and len(files) == 3
    print(json.dumps(dict(stage="engine_timing_complete", seconds=result["engine"]["wall_seconds"],
                          peak_rss_MiB=result["peak_rss_MiB"], run=str(run))), flush=True)
    result["oracle"] = verify_openpyxl(source, files, fixture["expected_row_sha256"])
    result["source_sha256_after_oracle"] = sha(source)
    assert result["source_sha256_after_oracle"] == result["source_sha256_after"] == before
    result["outputs"] = [dict(name=path.name, sha256=sha(path), bytes=path.stat().st_size) for path in files]
    result["passed"] = True
    result["scope"] = "Synthetic 100000-row large-SST input; engine process only measured; oracle and fixture creation excluded; no extrapolation to business workbooks."
    (run / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(dict(result=str(run / "result.json"), passed=True,
                          seconds=result["engine"]["wall_seconds"], peak_rss_MiB=result["peak_rss_MiB"])))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", nargs=3)
    options = parser.parse_args()
    if options.worker:
        worker(*(Path(value) for value in options.worker))
    else:
        main()
