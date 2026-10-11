"""Independent, synthetic XML fixture for escaped-string interpretation."""
import argparse
import hashlib
import inspect
import json
from pathlib import Path
import sys
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / 'excel_table_io.py').is_file())
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--workdir', type=Path, required=True, help='A new temporary or ignored directory; must not already exist.')
HERE = parser.parse_args().workdir.resolve()
HERE.mkdir(parents=True, exist_ok=False)
sys.path.insert(0, str(ROOT))
import openpyxl
from openpyxl.worksheet._reader import WorkSheetParser
from excel_table_io import TableReader, XlsxWriter

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
values = ['left_x003D_right', '_x003D_1+1', '_x005F_x003D_', 'left_x000D_right', 'left&#13;right', '_xD83D__xDE00_']
expected = ['left=right', '=1+1', '_x003D_', 'left\rright', 'left\rright', '\U0001f600']
row1 = ''.join(f'<c r="{chr(65+i)}1" t="inlineStr"><is><t>Case{i+1}</t></is></c>' for i in range(6))
row2 = ''.join(f'<c r="{chr(65+i)}2" t="inlineStr"><is><t>{v}</t></is></c>' for i,v in enumerate(values))
parts = {
    '[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
    '_rels/.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="wb" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
    'xl/workbook.xml': f'<workbook xmlns="{NS}" xmlns:r="{REL}"><sheets><sheet name="Synthetic" sheetId="1" r:id="s"/></sheets></workbook>',
    'xl/_rels/workbook.xml.rels': f'<Relationships xmlns="{PKG}"><Relationship Id="s" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
    'xl/worksheets/sheet1.xml': f'<worksheet xmlns="{NS}"><dimension ref="A1:F2"/><sheetData><row r="1">{row1}</row><row r="2">{row2}</row></sheetData></worksheet>',
}
source, output = HERE / 'synthetic-source.xlsx', HERE / 'synthetic-output.xlsx'
with source.open('xb') as stream, ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
    for name, body in parts.items():
        archive.writestr(name, body)
with TableReader(source, 'Synthetic') as reader, XlsxWriter(output, reader.headers) as writer:
    for _, cells in reader.iter_rows():
        writer.append(cells)
result = {'synthetic_only': True, 'openpyxl_version': openpyxl.__version__, 'expected': expected}
for kind, path in [('source', source), ('output', output)]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        cells = next(book.worksheets[0].iter_rows(min_row=2, max_row=2))
        result[kind] = [{'value': c.value, 'type': c.data_type} for c in cells]
    finally:
        book.close()
start = inspect.getsourcelines(WorkSheetParser.parse_cell)[1]
result['installed_openpyxl_parse_cell'] = {'module': 'openpyxl/worksheet/_reader.py', 'start_line': start}
result['application_sha256'] = hashlib.sha256((ROOT / 'excel_table_io.py').read_bytes()).hexdigest()
(HERE / 'openpyxl-result.json').write_text(json.dumps(result, indent=2, ensure_ascii=True), encoding='utf-8')
(HERE / 'expected.json').write_text(json.dumps(expected, ensure_ascii=True), encoding='utf-8')
print(json.dumps(result, indent=2, ensure_ascii=True))
