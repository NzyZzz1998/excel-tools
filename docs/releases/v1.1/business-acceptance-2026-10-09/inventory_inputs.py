"""Read-only, value-redacted input inventory; does not import the processing engine."""

import hashlib
import json
import posixpath
import sys
import warnings
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries


SOURCE_DIR = Path(r"E:\codex\excel-tools\testfile\期间缺货Top20")
SOURCE_NAMES = (
    "供应商缺货Top20 (1).xlsx",
    "库存满足率 (1).xlsx",
    "每日缺货Top50 (1).xlsx",
    "期间缺货Top20.xlsx",
    "缺货数据统计.xlsx",
)
REPORT_PATH = Path(__file__).resolve().with_name("input-inventory.json")


def digest(payload):
    return hashlib.sha256(payload).hexdigest()


def local(tag):
    return tag.rsplit("}", 1)[-1]


def elements(parent, name):
    return [child for child in parent if local(child.tag) == name]


def has_payload(cell):
    if cell is None:
        return False
    if elements(cell, "f"):
        return True
    return any(child.text for child in elements(cell, "v")) or any(
        child.text for child in cell.iter() if local(child.tag) == "t")


def metadata_reasons(cell):
    if cell is None:
        return []
    reasons = [name for name in ("vm", "cm") if name in cell.attrib]
    if elements(cell, "extLst"):
        reasons.append("extLst")
    return reasons


def normalized_part(base, target):
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(
        posixpath.join(posixpath.dirname(base), target))


def sheet_inventory(raw, sheet_number, name, state, path):
    root = ET.fromstring(raw)
    data = elements(root, "sheetData")[0]
    cells = {}
    raw_types, styles = Counter(), Counter()
    missing_rows = missing_cells = 0
    row_number = 0
    rows = elements(data, "row")
    for row in rows:
        if "r" not in row.attrib:
            missing_rows += 1
        row_number = int(row.attrib.get("r", row_number + 1))
        column = 0
        for cell in elements(row, "c"):
            if "r" not in cell.attrib:
                missing_cells += 1
                column += 1
                key = (row_number, column)
            else:
                key = coordinate_to_tuple(cell.attrib["r"])
                column = key[1]
            cells[key] = cell
            raw_types[cell.attrib.get("t", "n")] += 1
            styles[cell.attrib.get("s", "implicit")] += 1
    merges = []
    for group in elements(root, "mergeCells"):
        for merge in elements(group, "mergeCell"):
            ref = merge.attrib["ref"]
            first_col, first_row, last_col, last_row = range_boundaries(ref)
            height, width = last_row - first_row + 1, last_col - first_col + 1
            kind = ("rectangle" if height > 1 and width > 1 else
                    "vertical_only" if height > 1 else
                    "horizontal_only" if width > 1 else "single_cell")
            merges.append({"ref": ref, "r1": first_row, "c1": first_col,
                           "r2": last_row, "c2": last_col, "kind": kind})
    overlap_pairs = []
    for index, left in enumerate(merges):
        for right in merges[index + 1:]:
            if (max(left["r1"], right["r1"]) <= min(left["r2"], right["r2"])
                    and max(left["c1"], right["c1"]) <= min(left["c2"], right["c2"])):
                overlap_pairs.append([left["ref"], right["ref"]])

    def prediction(all_merges):
        chosen = [merge for merge in merges if merge["r2"] > merge["r1"]
                  or (all_merges and merge["c2"] > merge["c1"])]
        remaining = [merge["ref"] for merge in merges if merge not in chosen]
        blockers = []
        filled = 0
        anchor_types = Counter()
        for merge in chosen:
            r1, c1, r2, c2 = (merge[key] for key in ("r1", "c1", "r2", "c2"))
            anchor = cells.get((r1, c1))
            anchor_types["missing" if anchor is None else anchor.attrib.get("t", "n")] += 1
            if anchor is not None and elements(anchor, "f"):
                blockers.append({"kind": "formula_anchor", "merge": merge["ref"],
                                 "cell": anchor.attrib.get("r")})
            if metadata_reasons(anchor):
                blockers.append({"kind": "metadata_anchor", "merge": merge["ref"],
                                 "cell": anchor.attrib.get("r"),
                                 "metadata": metadata_reasons(anchor)})
            for (row, column), cell in cells.items():
                if r1 <= row <= r2 and c1 <= column <= c2 and (row, column) != (r1, c1):
                    if has_payload(cell) or metadata_reasons(cell):
                        blockers.append({"kind": "hidden_independent_content_or_metadata",
                                         "merge": merge["ref"], "cell": cell.attrib.get("r"),
                                         "has_formula": bool(elements(cell, "f")),
                                         "metadata": metadata_reasons(cell)})
            if has_payload(anchor):
                filled += ((r2 - r1 + 1) * (c2 - c1 + 1) - 1 if all_merges else r2 - r1)
            if not all_merges and c2 > c1:
                from openpyxl.utils.cell import get_column_letter
                remaining.extend("%s%d:%s%d" % (get_column_letter(c1), row,
                                                get_column_letter(c2), row)
                                 for row in range(r1, r2 + 1))
        if chosen and (missing_rows or missing_cells):
            blockers.append({"kind": "unsupported_omitted_coordinates",
                             "missing_row_r": missing_rows, "missing_cell_r": missing_cells})
        return {"selected_merge_count": len(chosen), "planned_filled_cell_count": filled,
                "remaining_merge_count": len(remaining), "remaining_merge_refs": remaining,
                "selected_anchor_xml_types": dict(anchor_types),
                "blockers": blockers, "overlapping_merge_pairs": overlap_pairs,
                "prediction_status": "blocked" if blockers else "needs_manual_interpretation"
                if overlap_pairs else "processable"}

    def refs_for(element):
        return [node.attrib.get("ref") or node.attrib.get("sqref")
                for node in elements(root, element)
                if node.attrib.get("ref") or node.attrib.get("sqref")]

    features = {feature: len(elements(root, feature)) for feature in (
        "autoFilter", "tableParts", "conditionalFormatting", "dataValidations", "drawing",
        "legacyDrawing", "hyperlinks", "sheetProtection", "extLst")}
    features["auto_filter_refs"] = refs_for("autoFilter")
    features["conditional_formatting_refs"] = refs_for("conditionalFormatting")
    declared_dimension = next((node.attrib.get("ref") for node in elements(root, "dimension")), None)
    actual_bounds = ({"min_row": min(row for row, _ in cells),
                      "max_row": max(row for row, _ in cells),
                      "min_column": min(column for _, column in cells),
                      "max_column": max(column for _, column in cells)} if cells else None)
    dimension_covers_cells = None
    if declared_dimension and cells:
        c1, r1, c2, r2 = range_boundaries(declared_dimension)
        dimension_covers_cells = (r1 <= actual_bounds["min_row"] <= actual_bounds["max_row"] <= r2
                                 and c1 <= actual_bounds["min_column"] <= actual_bounds["max_column"] <= c2)
    result = {
        "sheet_index": sheet_number, "sheet_name_sha256": digest(name.encode()),
        "sheet_state": state, "part": path, "xml_bytes": len(raw), "xml_sha256": digest(raw),
        "dimension_ref": declared_dimension,
        "actual_cell_bounds": actual_bounds,
        "declared_dimension_covers_actual_cells": dimension_covers_cells,
        "physical_rows": len(rows), "physical_cells": len(cells),
        "xml_cell_type_counts": dict(raw_types), "cell_style_index_counts": dict(styles),
        "custom_styled_row_count": sum("s" in row.attrib for row in rows),
        "hidden_row_count": sum(row.attrib.get("hidden") in ("1", "true") for row in rows),
        "custom_styled_column_definition_count": sum(
            "style" in col.attrib for group in elements(root, "cols") for col in elements(group, "col")),
        "missing_row_reference_count": missing_rows, "missing_cell_reference_count": missing_cells,
        "formula_cell_count": sum(bool(elements(cell, "f")) for cell in cells.values()),
        "metadata_cell_count": sum(bool(metadata_reasons(cell)) for cell in cells.values()),
        "merge_count": len(merges), "merge_kind_counts": dict(Counter(merge["kind"] for merge in merges)),
        "worksheet_features": features,
        "expected_default": prediction(False), "expected_all_merges": prediction(True),
    }
    return result


def inventory_file(path):
    source_bytes = path.read_bytes()
    result = {"file_name": path.name, "source_path": str(path), "source_bytes": len(source_bytes),
              "source_sha256_before": digest(source_bytes), "sheets": []}
    with ZipFile(path) as archive:
        package_names = archive.namelist()
        bad_part = archive.testzip()
        result["zip_integrity"] = "passed" if bad_part is None else "failed"
        result["zip_part_count"] = len(package_names)
        result["document_signature_part_count"] = sum(
            name.startswith("_xmlsignatures/") for name in package_names)
        result["part_category_counts"] = {
            "styles": sum(name.endswith("/styles.xml") for name in package_names),
            "shared_strings": sum(name.endswith("/sharedStrings.xml") for name in package_names),
            "charts": sum(name.startswith("xl/charts/") and name.endswith(".xml") for name in package_names),
            "drawings": sum(name.startswith("xl/drawings/") and name.endswith(".xml") for name in package_names),
            "media": sum(name.startswith("xl/media/") for name in package_names),
            "tables": sum(name.startswith("xl/tables/") and name.endswith(".xml") for name in package_names),
            "pivot_tables": sum(name.startswith("xl/pivotTables/") and name.endswith(".xml") for name in package_names),
            "external_links": sum(name.startswith("xl/externalLinks/") for name in package_names),
            "vba": sum("vba" in name.lower() for name in package_names),
            "metadata_or_rich_data": sum("metadata" in name.lower() or "richdata" in name.lower()
                                         for name in package_names),
        }
        package_rels = ET.fromstring(archive.read("_rels/.rels"))
        workbook_part = next(normalized_part("", node.attrib["Target"])
                             for node in package_rels if node.attrib.get("Type", "").endswith("/officeDocument"))
        workbook = ET.fromstring(archive.read(workbook_part))
        workbook_rel_part = posixpath.join(posixpath.dirname(workbook_part), "_rels",
                                          posixpath.basename(workbook_part) + ".rels")
        relationships = ET.fromstring(archive.read(workbook_rel_part))
        paths = {node.attrib["Id"]: normalized_part(workbook_part, node.attrib["Target"])
                 for node in relationships if node.attrib.get("Type", "").endswith("/worksheet")
                 and node.attrib.get("TargetMode") != "External"}
        result["workbook_part"] = workbook_part
        result["standard_workbook_path"] = workbook_part == "xl/workbook.xml"
        result["workbook_protected"] = bool(elements(workbook, "workbookProtection"))
        result["defined_name_count"] = sum(len(group) for group in elements(workbook, "definedNames"))
        worksheet_parts = set(paths.values())
        result["nonworksheet_part_sha256"] = {
            name: digest(archive.read(name)) for name in package_names if name not in worksheet_parts}
        if "xl/styles.xml" in package_names:
            styles = ET.fromstring(archive.read("xl/styles.xml"))
            result["style_definition_counts"] = {local(node.tag): len(node) for node in styles}
        if "xl/sharedStrings.xml" in package_names:
            strings = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            result["shared_string_count"] = len(strings)
        for group in elements(workbook, "sheets"):
            for sheet_number, node in enumerate(group, 1):
                relationship_id = next(value for key, value in node.attrib.items()
                                       if key.startswith("{") and local(key) == "id")
                if relationship_id not in paths:
                    continue
                part = paths[relationship_id]
                result["sheets"].append(sheet_inventory(archive.read(part), sheet_number,
                                                        node.attrib["name"], node.attrib.get("state", "visible"), part))
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter("always")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result["openpyxl_warning_codes"] = [
        "missing_default_cell_style" if "no default style" in str(item.message)
        else item.category.__name__ for item in captured]
    try:
        result["openpyxl_readable"] = True
        for index, sheet in enumerate(workbook.worksheets, 1):
            declared_openpyxl_shape = {"max_row": sheet.max_row, "max_column": sheet.max_column}
            # These source exports declare dimension=A1 despite containing many cells.
            # Ignore that hint and read the entire XML stream instead of clipping to A1.
            sheet.reset_dimensions()
            types = Counter()
            formula_count = 0
            nonempty = 0
            observed_max_row = observed_max_column = 0
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    observed_max_row = max(observed_max_row, cell.row)
                    observed_max_column = max(observed_max_column, cell.column)
                    nonempty += 1
                    if cell.data_type == "f":
                        types["formula"] += 1
                        formula_count += 1
                    elif cell.data_type == "e":
                        types["error"] += 1
                    elif cell.is_date:
                        types["date_or_time"] += 1
                    elif isinstance(value, bool):
                        types["boolean"] += 1
                    elif isinstance(value, (int, float)):
                        types["number"] += 1
                    else:
                        types["text"] += 1
            inventory = result["sheets"][index - 1]
            inventory["openpyxl_nonempty_cells"] = nonempty
            inventory["openpyxl_value_type_counts"] = dict(types)
            inventory["openpyxl_formula_count"] = formula_count
            inventory["openpyxl_declared_shape"] = declared_openpyxl_shape
            inventory["openpyxl_dimensions_reset"] = True
            inventory["openpyxl_nonempty_max_row"] = observed_max_row
            inventory["openpyxl_nonempty_max_column"] = observed_max_column
    finally:
        workbook.close()
    for mode in ("default", "all_merges"):
        prediction = [sheet["expected_" + mode] for sheet in result["sheets"]]
        blocked = result["document_signature_part_count"] or any(item["blockers"] for item in prediction)
        selected = sum(item["selected_merge_count"] for item in prediction)
        result["expected_" + mode] = {
            "file_outcome": "rejected" if blocked else "saved" if selected else "no_output_needed",
            "selected_merge_count": selected,
            "planned_filled_cell_count": sum(item["planned_filled_cell_count"] for item in prediction),
            "remaining_merge_count": sum(item["remaining_merge_count"] for item in prediction),
            "blocker_count": sum(len(item["blockers"]) for item in prediction),
        }
    result["source_sha256_after"] = digest(path.read_bytes())
    result["source_unchanged"] = result["source_sha256_before"] == result["source_sha256_after"]
    return result


def main():
    report = {
        "recorded_at": datetime.now(timezone(timedelta(hours=8))).isoformat(),
        "scope": str(SOURCE_DIR), "privacy": "No cell values, formula text, supplier or product values recorded.",
        "method": "Independent ZIP/ElementTree/Openpyxl reads; no processing engine imports or calls.",
        "python_version": sys.version.split()[0], "openpyxl_version": openpyxl.__version__,
        "script_sha256": digest(Path(__file__).read_bytes()),
        "files": [inventory_file(SOURCE_DIR / name) for name in SOURCE_NAMES],
    }
    sheets = [sheet for item in report["files"] for sheet in item["sheets"]]
    value_types = Counter()
    merge_types = Counter()
    for item in sheets:
        value_types.update(item["openpyxl_value_type_counts"])
        merge_types.update(item["merge_kind_counts"])
    report["summary"] = {
        "source_file_count": len(report["files"]), "sheet_count": len(sheets),
        "all_sources_unchanged": all(item["source_unchanged"] for item in report["files"]),
        "hidden_sheet_count": sum(item["sheet_state"] != "visible" for item in sheets),
        "physical_cell_count": sum(item["physical_cells"] for item in sheets),
        "nonempty_cell_count": sum(item["openpyxl_nonempty_cells"] for item in sheets),
        "value_type_counts": dict(value_types), "merge_kind_counts": dict(merge_types),
        "incorrect_dimension_sheet_count": sum(
            item["declared_dimension_covers_actual_cells"] is False for item in sheets),
        "missing_default_style_file_count": sum(
            "missing_default_cell_style" in item["openpyxl_warning_codes"] for item in report["files"]),
        "expected_modes": {mode: {
            "selected_merge_count": sum(item["expected_" + mode]["selected_merge_count"]
                                        for item in report["files"]),
            "planned_filled_cell_count": sum(item["expected_" + mode]["planned_filled_cell_count"]
                                             for item in report["files"]),
            "remaining_merge_count": sum(item["expected_" + mode]["remaining_merge_count"]
                                         for item in report["files"]),
            "blocker_count": sum(item["expected_" + mode]["blocker_count"] for item in report["files"]),
        } for mode in ("default", "all_merges")},
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(REPORT_PATH), "files": [
        {"file_name": item["file_name"], "sha256": item["source_sha256_before"],
         "sheet_count": len(item["sheets"]), "default": item["expected_default"],
         "all_merges": item["expected_all_merges"], "source_unchanged": item["source_unchanged"]}
        for item in report["files"]]}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
