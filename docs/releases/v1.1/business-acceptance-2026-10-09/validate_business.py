"""Local business acceptance; never writes into the supplied source directory.

Reads workbook content for comparison but records only counts, hashes and cell
addresses. Expected values are derived independently from source merge ranges.
"""
import argparse
import copy
import hashlib
import json
import posixpath
import shutil
import sys
import time
import warnings
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

REPO = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(REPO))
SOURCE = REPO / 'testfile' / '期间缺货Top20'
WORK = REPO / 'testfile' / '验收_v1.1_2026-10-09'
EVIDENCE = Path(__file__).resolve().parent


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def local(tag):
    return tag.rsplit('}', 1)[-1]


def signature(node):
    # xmlns declarations are resolved by ElementTree; lexical prefix changes
    # must not be mistaken for value or namespace changes.
    return (node.tag, tuple(sorted(node.attrib.items())), node.text,
            tuple((signature(child), child.tail) for child in node))


def cell_content(node):
    return (tuple(sorted((key, val) for key, val in node.attrib.items() if key != 'r')),
            tuple(signature(child) for child in node))


def workbook_parts(archive):
    relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    paths = {}
    for relation in relations:
        if relation.get('Type', '').endswith('/worksheet'):
            target = relation.get('Target')
            paths[relation.get('Id')] = (target.lstrip('/') if target.startswith('/')
                                       else posixpath.normpath('xl/' + target))
    book = ET.fromstring(archive.read('xl/workbook.xml'))
    result = []
    for node in book.iter():
        if local(node.tag) == 'sheet':
            relationship = next(v for k, v in node.attrib.items() if k.endswith('}id'))
            result.append((node.get('name'), paths[relationship]))
    return result


def inspect_sheet(raw):
    root = ET.fromstring(raw)
    data = next(node for node in root if local(node.tag) == 'sheetData')
    rows = {int(node.get('r')): node for node in data if local(node.tag) == 'row'}
    cells = {cell.get('r'): cell for row in rows.values() for cell in row if local(cell.tag) == 'c'}
    merges = [node.get('ref') for group in root if local(group.tag) == 'mergeCells' for node in group]
    return root, rows, cells, merges


def compare(source, output, all_merges):
    issues, sheet_results = [], []
    changed_parts = set()
    filled_total = regions_total = cells_checked = 0
    with ZipFile(source) as before, ZipFile(output) as after:
        if before.namelist() != after.namelist() or before.comment != after.comment:
            issues.append('archive entries/order/comment changed')
        for index, (name, part) in enumerate(workbook_parts(before), 1):
            aroot, arows, acells, merges = inspect_sheet(before.read(part))
            broot, brows, bcells, remaining = inspect_sheet(after.read(part))
            expected = {ref: cell_content(node) for ref, node in acells.items()}
            expected_merges, targets, changed_rows = [], {}, set()
            selected = 0
            for merged in merges:
                c1, r1, c2, r2 = range_boundaries(merged)
                if r1 == r2 and (not all_merges or c1 == c2):
                    expected_merges.append(merged)
                    continue
                selected += 1
                anchor_ref = '{}{}'.format(get_column_letter(c1), r1)
                anchor = acells.get(anchor_ref)
                attrs = tuple(sorted((k, v) for k, v in anchor.attrib.items() if k in ('s', 't'))) if anchor is not None else ()
                nodes = tuple(signature(child) for child in anchor if local(child.tag) in ('v', 'is')) if anchor is not None else ()
                has_value = anchor is not None and any(
                    node.text for node in anchor.iter() if local(node.tag) in ('v', 't'))
                for row in range(r1, r2 + 1):
                    changed_rows.add(row)
                    if not all_merges and c2 > c1:
                        expected_merges.append('{}{}:{}{}'.format(get_column_letter(c1), row, get_column_letter(c2), row))
                    for col in range(c1, (c2 if all_merges else c1) + 1):
                        ref = '{}{}'.format(get_column_letter(col), row)
                        if ref == anchor_ref:
                            continue
                        expected[ref] = (attrs, nodes)
                        targets[ref] = anchor_ref
                        filled_total += bool(has_value)
            if selected:
                changed_parts.add(part)
            elif before.read(part) != after.read(part):
                issues.append('sheet {} unchanged worksheet bytes differ'.format(index))
            actual = {ref: cell_content(node) for ref, node in bcells.items()}
            for ref in set(expected) | set(actual):
                if expected.get(ref) != actual.get(ref):
                    issues.append('sheet {} cell {} XML semantic mismatch'.format(index, ref))
            if sorted(expected_merges) != sorted(remaining):
                issues.append('sheet {} remaining merges differ'.format(index))
            dimension = next((n.get('ref') for n in broot if local(n.tag) == 'dimension'), None)
            if dimension is not None:
                c1, r1, c2, r2 = range_boundaries(dimension)
                for ref in list(bcells) + remaining:
                    a, b, c, d = range_boundaries(ref)
                    if not (c1 <= a <= c <= c2 and r1 <= b <= d <= r2):
                        issues.append('sheet {} dimension excludes {}'.format(index, ref))
                        break
            for row, node in arows.items():
                attrs = dict(node.attrib)
                if row in changed_rows:
                    attrs.pop('spans', None)
                if row not in brows or brows[row].attrib != attrs:
                    issues.append('sheet {} row {} attributes differ'.format(index, row))
            if not set(arows).issubset(brows):
                issues.append('sheet {} lost source rows'.format(index))
            if list(brows) != sorted(brows):
                issues.append('sheet {} row ordering invalid'.format(index))
            ignored = {'sheetData', 'mergeCells', 'dimension'}
            if aroot.attrib != broot.attrib or [signature(n) for n in aroot if local(n.tag) not in ignored] != [signature(n) for n in broot if local(n.tag) not in ignored]:
                issues.append('sheet {} non-target worksheet settings differ'.format(index))
            regions_total += selected
            cells_checked += len(expected)
            sheet_results.append({'sheet_index': index, 'part': part,
                                  'source_cells': len(acells), 'expected_output_cells': len(expected),
                                  'selected_regions': selected, 'target_count': len(targets),
                                  'remaining_merges': len(expected_merges)})
        for part in before.namelist():
            if part not in changed_parts and before.read(part) != after.read(part):
                issues.append('untouched package part changed: ' + part)

    # Independent reader verifies cell values, data types and number formats.
    # Business exports have inaccurate dimension metadata, so read_only=False.
    reader_warnings = []
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        original = openpyxl.load_workbook(source, read_only=False, data_only=False)
        transformed = openpyxl.load_workbook(output, read_only=False, data_only=False)
        streamed = openpyxl.load_workbook(output, read_only=True, data_only=False)
        try:
            for index, (src, dst) in enumerate(zip(original.worksheets, transformed.worksheets), 1):
                expected_values = {cell.coordinate: (cell.value, cell.data_type, cell.number_format)
                                   for row in src for cell in row if cell.value is not None}
                for merge in src.merged_cells.ranges:
                    c1, r1, c2, r2 = merge.bounds
                    if r1 == r2 and not all_merges:
                        continue
                    anchor = src.cell(r1, c1)
                    if anchor.value is None:
                        continue
                    for row in range(r1, r2 + 1):
                        for col in range(c1, (c2 if all_merges else c1) + 1):
                            expected_values['{}{}'.format(get_column_letter(col), row)] = (anchor.value, anchor.data_type, anchor.number_format)
                actual_values = {cell.coordinate: (cell.value, cell.data_type, cell.number_format)
                                 for row in dst for cell in row if cell.value is not None}
                for ref in set(expected_values) | set(actual_values):
                    if expected_values.get(ref) != actual_values.get(ref):
                        issues.append('sheet {} cell {} independent reader mismatch'.format(index, ref))
                streamed_values = {cell.coordinate: (cell.value, cell.data_type, cell.number_format)
                                   for row in streamed.worksheets[index - 1] for cell in row if cell.value is not None}
                if streamed_values != actual_values:
                    issues.append('sheet {} default streaming reader differs from full reader'.format(index))
        finally:
            original.close()
            transformed.close()
            streamed.close()
        reader_warnings = [str(item.message) for item in caught]
    return {'passed': not issues, 'issues': issues, 'sheet_results': sheet_results,
            'selected_regions': regions_total, 'filled_cells': filled_total,
            'xml_cells_checked': cells_checked, 'reader_warnings': reader_warnings,
            'output_sha256': sha(output)}


def run():
    parser = argparse.ArgumentParser()
    parser.add_argument('--validate-exe', action='store_true')
    parser.add_argument('--revision', default='')
    args = parser.parse_args()
    sources = sorted(SOURCE.glob('*.xlsx'))
    identity = {name: sha(REPO / name) for name in (
        'excel_unmerge_fill.py', 'excel_unmerge_gui.py', 'dist/ExcelTools-v1.1-Windows-x64.zip')}
    hashes = {p.name: sha(p) for p in sources}
    results = []
    for mode in ('default', 'all'):
        dest = WORK / args.revision / ('exe' if args.validate_exe else 'engine') / mode
        dest.mkdir(parents=True, exist_ok=True)
        for original in sources:
            source = dest / original.name
            if args.validate_exe:
                outputs = sorted(dest.glob(original.stem + '_拆分填充*.xlsx'))
                if len(outputs) != 1:
                    results.append({'file': original.name, 'mode': mode, 'passed': False,
                                    'issues': ['expected exactly one EXE output; found ' + str(len(outputs))]})
                    continue
                output, stats, elapsed = outputs[0], None, None
            else:
                if source.exists():
                    raise RuntimeError('Work directory already used; preserve previous evidence: ' + str(dest))
                shutil.copy2(original, source)
                from excel_unmerge_fill import process_file
                start = time.perf_counter()
                output, stats = process_file(source, all_merges=mode == 'all')
                elapsed = time.perf_counter() - start
                if output is None:
                    results.append({'file': original.name, 'mode': mode, 'no_output': True,
                                    'stats': stats, 'passed': False,
                                    'issues': ['all supplied fixtures require an output']})
                    continue
            result = compare(source, output, mode == 'all')
            result.update({'file': original.name, 'mode': mode, 'source': str(source),
                           'source_sha256': sha(source), 'output': str(output), 'seconds': elapsed,
                           'source_matches_original': sha(source) == hashes[original.name]})
            if stats is not None:
                result['reported_stats_match'] = (sum(x[1] for x in stats) == result['selected_regions']
                                                  and sum(x[2] for x in stats) == result['filled_cells'])
                result['passed'] = result['passed'] and result['reported_stats_match']
            result['passed'] = result['passed'] and result['source_matches_original']
            results.append(result)
    current = {p.name: sha(p) for p in sources}
    report = {'version': '1.1', 'track': 'exe' if args.validate_exe else 'source-engine',
              'revision': args.revision, 'candidate_identity': identity,
              'validation_script_sha256': sha(Path(__file__)),
              'default_streaming_reader_checked': True,
              'source_dir': str(SOURCE), 'original_hashes_before': hashes,
              'original_hashes_after': current, 'all_originals_unchanged': current == hashes,
              'results': results, 'passed': all(r['passed'] for r in results) and current == hashes}
    suffix = '-' + args.revision if args.revision else ''
    path = EVIDENCE / (('exe' if args.validate_exe else 'engine') + '-content-validation' + suffix + '.json')
    if path.exists():
        raise RuntimeError('Preserve existing evidence: ' + str(path))
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'report': str(path), 'passed': report['passed'], 'results': [
        {k: r.get(k) for k in ('file', 'mode', 'passed', 'selected_regions', 'filled_cells', 'xml_cells_checked', 'issues')}
        for r in results]}, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    run()
