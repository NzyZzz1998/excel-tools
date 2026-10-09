"""Release GUI boundary probes; only synthetic workbooks and owned transparent Tk windows."""
import errno
import hashlib
import json
import queue
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch
from xml.parsers.expat import ExpatError
from zipfile import BadZipFile

from openpyxl import Workbook

ROOT = Path(sys.argv.pop(1)).resolve()
sys.path.insert(0, str(ROOT))
import excel_unmerge_gui as gui
from excel_unmerge_fill import process_file


class GuiBoundaryProbes(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="excel-release-gui-")
        self.folder = Path(self.directory.name)
        self.root = tk.Tk()
        self.root.attributes("-alpha", 0)
        self.callback_errors = []
        self.root.report_callback_exception = lambda *args: self.callback_errors.append(str(args))
        self.app = gui.Application(self.root)
        self.root.update()

    def tearDown(self):
        if self.root.winfo_exists():
            self.root.destroy()
        self.directory.cleanup()
        self.assertEqual(self.callback_errors, [])

    def workbook(self, name):
        path = self.folder / name
        book = Workbook()
        book.active["A1"] = "synthetic group"
        book.active.merge_cells("A1:A3")
        book.save(path)
        book.close()
        return str(path)

    def finish(self):
        deadline = time.monotonic() + 5
        while self.app.running and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.005)
        self.assertFalse(self.app.running)

    def test_thread_start_failure_keeps_pending_files_then_recovers(self):
        source = self.workbook("thread-start.xlsx")
        self.app.files = (source,)
        with patch.object(threading.Thread, "start", side_effect=RuntimeError("synthetic thread limit")):
            self.app.start()
            self.finish()
        self.assertEqual(self.app.pending_files, (source,))
        self.assertEqual((self.app.completed, self.app.succeeded, self.app.failed), (0, 0, 0))
        self.assertNotIn("disabled", self.app.resume_button.state())
        self.assertIn("synthetic thread limit", self.app.results.get("1.0", "end"))
        self.app.resume_pending()
        self.finish()
        self.assertEqual(self.app.succeeded, 1)
        self.assertEqual(self.app.pending_files, ())

    def test_interrupted_retry_and_resume_preserve_disjoint_recovery_sets(self):
        a, b, c, d = [self.workbook(name + ".xlsx") for name in "abcd"]
        self.app.files = (a, b, c, d)

        def first_attempt(file, all_merges):
            if file == b:
                self.app.stop_event.set()
            raise PermissionError("synthetic initial failure")

        with patch.object(gui, "process_file", side_effect=first_attempt):
            self.app.start()
            self.finish()
        self.assertEqual(self.app.failed_files, (a, b))
        self.assertEqual(self.app.pending_files, (c, d))

        def save_then_stop(file, all_merges):
            result = process_file(file, all_merges=all_merges)
            self.app.stop_event.set()
            return result

        with patch.object(gui, "process_file", side_effect=save_then_stop):
            self.app.retry_failed()
            self.finish()
        self.assertEqual(self.app.failed_files, (b,))
        self.assertEqual(self.app.pending_files, (c, d))
        with patch.object(gui, "process_file", side_effect=save_then_stop):
            self.app.resume_pending()
            self.finish()
        self.assertEqual(self.app.failed_files, (b,))
        self.assertEqual(self.app.pending_files, (d,))
        self.assertNotIn("disabled", self.app.retry_button.state())
        self.assertNotIn("disabled", self.app.resume_button.state())
        self.app.retry_failed()
        self.finish()
        self.assertEqual(self.app.failed_files, ())
        self.assertEqual(self.app.pending_files, (d,))
        self.app.resume_pending()
        self.finish()
        self.assertEqual((self.app.completed, self.app.succeeded, self.app.failed), (4, 4, 0))
        self.assertEqual(len(list(self.folder.glob("*_拆分填充.xlsx"))), 4)
        self.assertEqual(list(self.folder.glob("*_拆分填充_2.xlsx")), [])

    def test_late_stop_after_worker_done_does_not_invent_pending_files(self):
        source = self.workbook("late-stop.xlsx")
        self.app.files = (source,)
        self.app.start()
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            with self.app.events.mutex:
                done = any(event[0] == "done" for event in self.app.events.queue)
            if done:
                break
            time.sleep(.005)
        self.assertTrue(done)
        self.app.request_stop()
        self.finish()
        self.assertEqual((self.app.succeeded, self.app.completed), (1, 1))
        self.assertEqual(self.app.pending_files, ())
        self.assertNotIn("已停止", self.app.status.get())

    def test_repeated_close_and_actions_during_active_file_do_not_reenter(self):
        first = self.workbook("active.xlsx")
        later = self.workbook("later.xlsx")
        entered, release = threading.Event(), threading.Event()

        def delayed(file, all_merges):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Synthetic release timed out")
            return process_file(file, all_merges=all_merges)

        self.app.files = (first, later)
        with patch.object(gui, "process_file", side_effect=delayed) as calls:
            self.app.start()
            try:
                self.assertTrue(entered.wait(5))
                self.app.start()
                self.app.retry_failed()
                self.app.resume_pending()
                self.app.close()
                self.app.close()
                self.app.request_stop()
                self.assertTrue(self.root.winfo_exists())
                self.assertTrue(self.app.running)
            finally:
                release.set()
                self.finish()
            self.assertEqual(calls.call_count, 1)
        self.assertEqual(self.app.pending_files, (later,))
        self.assertEqual(self.app.succeeded, 1)

    def test_choose_cancel_and_new_selection_follow_existing_task_contract(self):
        old = str(self.folder / "old-missing.xlsx")
        new = self.workbook("new-selection.xlsx")
        self.app.files = (old,)
        self.app.start()
        self.finish()
        old_log = self.app.results.get("1.0", "end")
        with patch.object(gui.filedialog, "askopenfilenames", return_value=()):
            self.app.choose()
        self.assertEqual(self.app.files, (old,))
        self.assertEqual(self.app.results.get("1.0", "end"), old_log)
        with patch.object(gui.filedialog, "askopenfilenames", return_value=(new,)):
            self.app.choose()
        self.assertEqual(self.app.failed_files, (old,))
        self.app.retry_failed()
        self.finish()
        self.assertEqual(self.app.failed_files, (old,))
        self.assertFalse(Path(new).with_name("new-selection_拆分填充.xlsx").exists())
        self.app.start()
        self.finish()
        self.assertEqual(self.app.failed_files, ())
        self.assertEqual(tuple(self.app.file_states), (new,))
        self.assertNotIn(old, self.app.results.get("1.0", "end"))
        self.assertEqual(self.app.succeeded, 1)

    def test_each_error_keeps_loop_live_and_all_failed_files_retryable(self):
        errors = [BadZipFile("synthetic zip"), PermissionError("synthetic permission"),
                  ExpatError("synthetic xml"), OSError(errno.ENOSPC, "synthetic full disk"),
                  MemoryError(), FileNotFoundError("synthetic missing")]
        files = tuple(str(self.folder / (str(i) + ".xlsx")) for i in range(len(errors)))
        self.app.files = files
        with patch.object(gui, "process_file", side_effect=errors):
            self.app.start()
            self.finish()
        self.assertEqual(self.app.failed_files, files)
        self.assertEqual(self.app.pending_files, ())
        self.assertEqual(self.app.failed, len(errors))
        self.assertEqual(self.app.completed, len(errors))
        self.assertIn("disabled", self.app.open_button.state())
        log = self.app.results.get("1.0", "end")
        for error in errors:
            self.assertIn(type(error).__name__, log)

    def test_moved_output_directory_reports_error_without_opening_external_window(self):
        source = self.workbook("output-folder.xlsx")
        self.app.files = (source,)
        self.app.start()
        self.finish()
        self.app.output_directory = self.folder / "no-longer-exists"
        with patch.object(gui.os, "startfile", create=True) as launch:
            self.app.open_results()
        launch.assert_not_called()
        self.assertIn("无法打开结果目录", self.app.status.get())
        self.assertIn("复制保存路径", self.app.status.get())


def long_name_observation():
    root = tk.Tk()
    root.attributes("-alpha", 0)
    try:
        app = gui.Application(root)
        root.geometry("680x540")
        with tempfile.TemporaryDirectory(prefix="excel-long-name-") as folder:
            source = Path(folder) / ("文" * 245 + ".xlsx")
            workbook = Workbook()
            workbook.active["A1"] = "synthetic value"
            workbook.active.merge_cells("A1:A3")
            workbook.save(source)
            workbook.close()
            entered, release = threading.Event(), threading.Event()

            def delayed(file, all_merges):
                entered.set()
                if not release.wait(5):
                    raise RuntimeError("Synthetic release timed out")
                return process_file(file, all_merges=all_merges)

            app.files = (str(source),)
            with patch.object(gui, "process_file", side_effect=delayed):
                app.start()
                try:
                    assert entered.wait(5)
                    deadline = time.monotonic() + 5
                    while app.current_file is None and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.005)
                    assert app.current_file == str(source)
                    root.update()
                    running_visible = bool(app.results.winfo_ismapped())
                    status_height = app.status_label.winfo_height()
                    stop_visible = bool(app.stop_button.winfo_ismapped())
                finally:
                    release.set()
                    deadline = time.monotonic() + 5
                    while app.running and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.005)
                    assert not app.running
            assert app.succeeded == 1
            return {"synthetic_workbook_executed": True, "window": "680x540", "filename_characters": 250,
                "running_status_height_px": status_height,
                "result_area_visible_while_long_status": running_visible,
                "stop_control_visible_while_long_status": stop_visible,
                "result_area_visible_after_summary": bool(app.results.winfo_ismapped()),
                "file_processed_successfully": app.succeeded == 1,
                "interpretation": "Nonblocking layout limitation at near-maximum filename length; stop controls remain available and completed results reappear."}
    finally:
        root.destroy()


if __name__ == "__main__":
    source = ROOT / "excel_unmerge_gui.py"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(GuiBoundaryProbes)
    outcome = unittest.TextTestRunner(verbosity=2).run(suite)
    report = {"source_sha256": before, "tests_run": outcome.testsRun,
              "tests_passed": outcome.wasSuccessful(),
              "source_unchanged": hashlib.sha256(source.read_bytes()).hexdigest() == before,
              "long_filename_layout": long_name_observation()}
    print("GUI_BOUNDARY_REPORT=" + json.dumps(report, ensure_ascii=False))
    raise SystemExit(0 if outcome.wasSuccessful() else 1)
